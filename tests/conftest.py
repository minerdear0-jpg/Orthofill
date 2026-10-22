import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
ORDER = ROOT / "order.xlsx"


@pytest.fixture
def order_xlsx(tmp_path) -> Path:
    """Копия order.xlsx во временной папке — оригинал тесты не трогают."""
    dst = tmp_path / "order.xlsx"
    shutil.copy(ORDER, dst)
    return dst


def _insert_rows(path: Path, rows) -> None:
    """Вставить пункты (код, текст) в A/B после последней строки их группы/раздела."""
    import openpyxl
    from orthofill.workbook import clean_code

    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    for code, text in rows:
        prefix = code[:2] if len(code) > 2 else code[0]
        last = max(c.row for (c,) in ws.iter_rows(min_col=1, max_col=1)
                   if c.value is not None and clean_code(c.value).startswith(prefix))
        ws.insert_rows(last + 1)
        ws.cell(last + 1, 1, code)
        ws.cell(last + 1, 2, text)
    wb.save(path)


@pytest.fixture
def order_xlsx_with(order_xlsx):
    """order_xlsx_with(("4f7", "Слон"), ...) -> Workbook с добавленными пунктами."""
    from orthofill.workbook import load_workbook_file

    def make(*rows):
        _insert_rows(order_xlsx, rows)
        return load_workbook_file(order_xlsx)
    return make


@pytest.fixture
def store(tmp_path):
    from orthofill.journal import Journal
    from orthofill.presets import PresetStore
    return PresetStore(Journal(tmp_path / "j.db"))


@pytest.fixture
def clip():
    return []


@pytest.fixture
def ctl(tmp_path, clip):
    from orthofill.controller import Controller
    from orthofill.journal import Journal
    return Controller(Journal(tmp_path / "j.db"), clip.append)


def _tk():
    """tk.Tk() с повтором: на Windows создание окна изредка не может прочитать свой
    файл Tcl (временная блокировка — например, антивирус при частом открытии)."""
    import time
    import tkinter as tk
    for attempt in range(3):
        try:
            return tk.Tk()
        except tk.TclError:
            if attempt == 2:
                raise
            time.sleep(0.5)


def _make_app(tmp_path, *paths):
    import tkinter as tk
    from orthofill.controller import Controller
    from orthofill.journal import Journal
    from orthofill.ui import App
    root = _tk()
    ctl = Controller(Journal(tmp_path / "j.db"), clipboard=lambda t: None)
    a = App(root, ctl)
    # живой протокол — только в своих тестах (test_live_ui задаёт задержку сам):
    # иначе черновик через 300 мс перезаписывает статус, который проверяет тест
    a.live_delay = a.unknown_delay = 600_000
    ctl.clipboard = a._clipboard
    for p in paths:
        ctl.import_xlsx(p)
    a._show_file_status()
    return a


@pytest.fixture
def app(tmp_path, order_xlsx):
    """Окно на настоящем Tk: пресеты «order» и «консультация» (текущий)."""
    other = tmp_path / "консультация.xlsx"
    shutil.copy(order_xlsx, other)
    a = _make_app(tmp_path, order_xlsx, other)
    yield a
    a.root.destroy()


@pytest.fixture
def app_with_p(tmp_path, order_xlsx_with):
    """Окно с пресетом, где в группе 4p челюсти перепутаны (как у врача)."""
    wb = order_xlsx_with(('4p', 'Эластическая цепочка'), ('4p1', 'на нижней челюсти.'),
                         ('4p2', 'на верхней челюсти.'))
    a = _make_app(tmp_path, wb.path)
    yield a
    a.root.destroy()


@pytest.fixture(autouse=True)
def _no_real_telemetry(monkeypatch):
    """Тесты никогда не ищут настоящий telemetry.json (он может лежать в папке
    проекта для сборки) — свой конфиг задают только тесты телеметрии."""
    import orthofill.controller as c
    monkeypatch.setattr(c, 'find_telemetry_config', lambda: None)
