"""0.3 beta 3: фразы пресета — хранение, xlsx, черновик."""
import sqlite3

import openpyxl
import pytest

from orthofill.phrases import parse_template
from orthofill.workbook import Phrase, load_workbook_file
from orthofill.xlsx_export import export_workbook

TPL = 'замена дуги {верх|низ|обе} {18х25|16х25} тма'


def test_parse_template():
    assert parse_template(TPL) == ['замена дуги ', ['верх', 'низ', 'обе'], ' ', ['18х25', '16х25'], ' тма']
    with pytest.raises(ValueError, match='не закрыта'):
        parse_template('замена {верх|низ')
    with pytest.raises(ValueError, match='пустое поле'):
        parse_template('замена {}')


def test_store_roundtrip(store, order_xlsx):
    wb = load_workbook_file(order_xlsx)
    wb.phrases = [Phrase('замена', TPL, 3)]
    pid = store.create(wb)
    back = store.load(pid)
    assert back.phrases == [Phrase('замена', TPL, 3)]


def test_xlsx_roundtrip(tmp_path, order_xlsx):
    wb = load_workbook_file(order_xlsx)
    wb.phrases = [Phrase('замена', TPL, 3), Phrase('акт', 'активация, {лигатурная|самолиг}', 0)]
    export_workbook(wb, tmp_path / 'x.xlsx')
    back = load_workbook_file(tmp_path / 'x.xlsx')
    assert back.phrases == wb.phrases and not back.problems


def test_bad_template_names_cell(tmp_path, order_xlsx):
    wb = load_workbook_file(order_xlsx)
    wb.phrases = [Phrase('замена', TPL, 0)]
    export_workbook(wb, tmp_path / 'x.xlsx')
    xl = openpyxl.load_workbook(tmp_path / 'x.xlsx')
    xl['Фразы'].append(['плохая', 'замена {верх|низ', 0])
    xl.save(tmp_path / 'x.xlsx')
    back = load_workbook_file(tmp_path / 'x.xlsx')
    assert any(p.startswith('Фразы!B3: не закрыта') for p in back.problems)


def test_migrate_v3_to_v4(tmp_path):
    from orthofill.journal import Journal, _SCHEMA, _SCHEMA_V2, _SCHEMA_V3
    db = tmp_path / 'v3.db'
    con = sqlite3.connect(db)
    con.executescript(_SCHEMA + _SCHEMA_V2 + _SCHEMA_V3)
    con.execute('ALTER TABLE entries ADD COLUMN preset_id INTEGER')
    con.execute('ALTER TABLE preset_items ADD COLUMN combine TEXT')
    con.execute('ALTER TABLE preset_items ADD COLUMN condition_json TEXT')
    con.execute('PRAGMA user_version = 3')
    con.commit()
    con.close()
    j = Journal(db)
    assert j._db.execute('PRAGMA user_version').fetchone()[0] == 4
    j._db.execute('SELECT name, template, uses FROM phrases')


def test_draft_keeps_phrases(ctl, order_xlsx):
    ctl.import_xlsx(order_xlsx)
    d = ctl.open_draft()
    d.wb.phrases.append(Phrase('замена', TPL, 0))
    d.set_text('4a', 'Осмотр.')
    ctl.save_draft()
    assert ctl.workbook.phrases == [Phrase('замена', TPL, 0)]
