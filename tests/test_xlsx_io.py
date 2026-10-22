"""Экспорт пресета в xlsx и обратное чтение: A–F, «Настройки», «Словарь»."""
import openpyxl
import pytest

from orthofill.engine import guess_settings
from orthofill.workbook import DictEntry, WorkbookError, load_workbook_file
from orthofill.xlsx_export import export_workbook


def props(wb):
    return [(i.id, i.text, i.method, i.when, i.default, i.words) for i in wb.items]


def refs(wb):
    return [(r.case_id, r.raw, r.lines) for r in wb.references]


def test_roundtrip(order_xlsx_with, tmp_path):
    wb = order_xlsx_with(("4o", "Подвязывание зуба короткой металлической лигатурой."))
    guess_settings(wb)
    wb.by_id["5d"].default = "5d3"
    wb.by_id["4b"].words = ["дуг", "проволок"]
    wb.dictionary = [DictEntry("торк", "Изгиб третьего порядка", "4g3", "2026-10-21", 2)]
    out = tmp_path / "out.xlsx"
    export_workbook(wb, out)
    wb2 = load_workbook_file(out)
    assert not wb2.problems
    assert props(wb2) == props(wb)
    assert refs(wb2) == refs(wb)
    assert wb2.dictionary == wb.dictionary


def test_file_01_has_empty_settings(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    assert wb.by_id["3f"].method is None and wb.dictionary == []


def test_unknown_method_label_is_problem(order_xlsx, tmp_path):
    wb = load_workbook_file(order_xlsx)
    out = tmp_path / "out.xlsx"
    export_workbook(wb, out)
    xl = openpyxl.load_workbook(out)
    ws = xl["Настройки"]
    ws.append(["3f", "по звёздам", None, None, None])
    xl.save(out)
    assert any("по звёздам" in p for p in load_workbook_file(out).problems)


def _sheet(tmp_path, rows):
    p = tmp_path / "x.xlsx"
    xl = openpyxl.Workbook()
    for r in rows:
        xl.active.append(r)
    xl.save(p)
    return p


def test_header_in_wrong_column(tmp_path):
    p = _sheet(tmp_path, [(None, None, "ID_CLASS", "CLASS_LIST"), (None, None, "1", "Жалобы")])
    with pytest.raises(WorkbookError, match="колонке C, а должна быть в A"):
        load_workbook_file(p)


def test_header_missing_message(tmp_path):
    p = _sheet(tmp_path, [(None, "Calibri"), ("1", "Жалобы")])
    with pytest.raises(WorkbookError, match="в колонке A должно стоять ID_CLASS"):
        load_workbook_file(p)


# ---------- 0.2 beta 5: логика — столбцы F, G, H ----------
def _logic_wb(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    wb.by_id["4f"].combine = "and"
    wb.by_id["4h1"].when = "condition"
    wb.by_id["4h1"].condition = {"all": [{"field": "2b", "is": "2b2"}]}
    wb.by_id["4a"].condition = {"any": [{"field": "4e", "mentioned": True},
                                        {"field": "2b", "is": "2b1", "not": True}]}
    wb.not_together = [("4h1", "4h2")]
    return wb


def test_logic_survives_export_import(tmp_path, order_xlsx):
    wb = _logic_wb(order_xlsx)
    export_workbook(wb, tmp_path / "x.xlsx")
    back = load_workbook_file(tmp_path / "x.xlsx")
    assert not back.problems and not back.warnings
    assert back.by_id["4f"].combine == "and" and back.by_id["4h1"].when == "condition"
    assert back.by_id["4h1"].condition == wb.by_id["4h1"].condition
    assert back.by_id["4a"].condition == wb.by_id["4a"].condition
    assert back.not_together == [("4h1", "4h2")]
    st = openpyxl.load_workbook(tmp_path / "x.xlsx")["Настройки"]
    assert [c.value for c in st[1]][5:8] == ["Сочетание", "Условие", "Не вместе с"]
    rows = {r[0]: r for r in st.iter_rows(min_row=2, values_only=True)}
    assert rows["4h1"][6] == "2b = 2b2" and rows["4h1"][7] == "4h2" and rows["4h2"][7] == "4h1"
    assert rows["4f"][5] == "И"


def _settings_rows(tmp_path, order_xlsx, rows):
    out = tmp_path / "s.xlsx"
    export_workbook(load_workbook_file(order_xlsx), out)
    xl = openpyxl.load_workbook(out)
    ws = xl["Настройки"]
    while ws.max_row < 11:
        ws.append([None])
    for r in rows:
        ws.append(r)
    xl.save(out)
    return load_workbook_file(out)


def test_condition_error_names_cell(tmp_path, order_xlsx):
    wb = _settings_rows(tmp_path, order_xlsx, [["4h1", None, "по условию", None, None, None, "4x = 4x1"]])
    assert any(p.startswith("Настройки!G12: не знаю код 4x") for p in wb.problems)


def test_bad_combine_and_pair_codes(tmp_path, order_xlsx):
    wb = _settings_rows(tmp_path, order_xlsx, [["4f", None, None, None, None, "ВСЕ", None, "9z"]])
    assert any(p.startswith("Настройки!F12") for p in wb.problems)
    assert any(p.startswith("Настройки!H12") and "9z" in p for p in wb.problems)


def test_asymmetric_pair_completed(tmp_path, order_xlsx):
    wb = _settings_rows(tmp_path, order_xlsx, [["4h1", None, None, None, None, None, None, "4h2"]])
    assert wb.not_together == [("4h1", "4h2")] and not wb.problems
    assert any("дополнена" in w for w in wb.warnings)
