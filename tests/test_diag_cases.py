"""Ошибки, найденные по диагностике врача (21.10.2026, сборка e7a33ad):
разбор записи и переименованная группа. Проверки протокола — данными,
в tests/cases/diag-2026-10-21.json.

Записи — дословно из diag-2026-10-21.jsonl. Пункты 4o/4p/4l врач добавил
в свой классификатор (order-new.xlsx) — здесь добавляем их так же."""
import pytest

from orthofill.engine import build_protocol
from orthofill.parse import extract_facts

DOCTOR_ROWS = [
    ("4o", "Подвязывание зуба короткой металлической лигатурой."),
    ("4p", "Эластическая цепочка без шага"),
    ("4p1", " на нижней челюсти."),
    ("4p2", "на верхней челюсти."),
    ("4l", "Активация тяги"),
    ("4l1", "к минивинту в 1 сегменте."),
    ("4l2", "к минивинту в 2 сегменте."),
    ("4l3", "к минивинту в 3 сегменте."),
    ("4l4", "к минивинту в 4 сегменте."),
]


@pytest.fixture
def wb(order_xlsx_with):
    return order_xlsx_with(*DOCTOR_ROWS)


def ids(p):
    return [c for c, _ in p.lines]


# ---------- челюсть цепочки/лигатур не переходит на дугу ----------
@pytest.mark.parametrize("raw, expected", [
    ("активация, система самолигируемая, гигиена удовлетворительная, класс 1 по молярам, "
     "цепочка на обеих челюстях, замена дуги 18 * 25 TMA  на верхней",
     {"upper": (("18", "25"), "TMA")}),
    ("эластическая цепочка на верхней, лигатуры на нижней,  замена дуги на верху 18 25 нити. "
     "ношение эластиков перепел",
     {"upper": (("18", "25"), "NiTi")}),
    # прежнее поведение: челюсть для дуги по-прежнему понимается
    ("лигатуры на нижней, верх 16х25 нити", {"upper": (("16", "25"), "NiTi")}),
    ("лигатурная дуга верх 16х25 нити явка 4 нед", {"upper": (("16", "25"), "NiTi")}),
])
def test_other_items_arch_not_used_for_wire(raw, expected):
    assert extract_facts(raw)["wires"] == expected


def test_ligature_system_word_is_not_ligature_change():
    assert extract_facts("самолигируемая, система лигатур")["ligature_change"] is False or \
        extract_facts("самолигируемая, система лигатур")["system_type"] == "self_ligating"


# ---------- опечатки в типе системы ----------
@pytest.mark.parametrize("raw", ["система самолинируемая", "система самолиuируемая",
                                 "самолиг", "самолигирующая"])
def test_self_ligating_typos(raw):
    assert extract_facts(raw)["system_type"] == "self_ligating"


# ---------- врач переименовал группу «Смена дуги на верхней челюсти» → «Смена дуги» ----------
def test_wire_group_arch_from_variants(order_xlsx):
    import openpyxl
    from orthofill.workbook import load_workbook_file
    from orthofill.engine import rule_warnings
    xl = openpyxl.load_workbook(order_xlsx)
    ws = xl.worksheets[0]
    for row in ws.iter_rows(min_col=1, max_col=2):
        if row[0].value == "4b":
            row[1].value = "Смена дуги"
    xl.save(order_xlsx)
    wb = load_workbook_file(order_xlsx)
    assert not any("дуга на верхней" in w for w in rule_warnings(wb))
    p = build_protocol(wb, "лигатурная, замена дуги верх 18х25 нити")
    assert "4b7" in ids(p) and "4c7" not in ids(p)