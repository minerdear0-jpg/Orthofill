"""Эталоны в новом виде: перевод старых, сверка по кодам, склеенные строки."""
from orthofill.engine import build_protocol, check_references
from orthofill.workbook import lines_for_reference, load_workbook_file, normalize_references


def test_order_reference_passes_after_normalize(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    assert all(r.ok for r in check_references(wb))
    lines = wb.references[0].lines
    assert ("2b", "Тип системы") in lines and ("2b2", "вестибулярная лигатурная.") in lines
    i = lines.index(("3f1", "Гигиена полости рта удовлетворительная."))
    assert lines[i - 1][0] == "3f"


def test_normalize_idempotent(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    before = [r.lines[:] for r in wb.references]
    normalize_references(wb)
    assert [r.lines for r in wb.references] == before


def test_normalize_unknown_join_kept(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    case = wb.references[0]
    i = case.lines.index(("2b", "Тип системы"))
    case.lines[i:i + 2] = [("2b", "Тип системы керамическая.")]
    normalize_references(wb)
    assert ("2b", "Тип системы керамическая.") in case.lines
    r = check_references(wb)[0]          # формулировка врача: не проверяется, но учтена
    assert r.ok and r.own_lines == 1


def test_saved_reference_roundtrip(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    p = build_protocol(wb, "лигатурная, перепел, явка 6 нед")
    assert lines_for_reference(wb, p, p.text) == p.lines


def test_merged_text_of_other_variant_maps_to_codes(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    p = build_protocol(wb, "лигатурная, явка 6 нед")
    text = p.text.replace("через 6 недель.", "через 8 недель.")
    lines = lines_for_reference(wb, p, text)
    assert ("5d", "Явка на следующую активацию через") in lines and ("5d4", "8 недель.") in lines


def test_edited_merged_row_is_own_line(ctl, order_xlsx):
    ctl.import_xlsx(order_xlsx)
    text = ctl.generate("лигатурная, явка 6 нед").text.replace("через 6 недель.", "через 5 недель.")
    ctl.save_reference("лигатурная, явка 6 нед", text)
    # формулировка врача для пункта «явка»: код группы, свой текст
    assert ("5d", "Явка на следующую активацию через 5 недель.") in ctl.workbook.references[-1].lines
    assert ctl.file_status == "✓ пунктов: 124 · эталоны 2/2"
    assert any("свои строки врача — 1" in d for d in ctl.details)


def test_preset_db_references_normalized(store, order_xlsx):
    pid = store.create(load_workbook_file(order_xlsx))
    store.add_reference(pid, "лигатурная", [("2", "Анамнез: активация"),
                                            ("2b", "Тип системы вестибулярная лигатурная.")])
    lines = store.load(pid).references[-1].lines
    assert lines == [("2", "Анамнез: активация"), ("2b", "Тип системы"),
                     ("2b2", "вестибулярная лигатурная.")]


def test_doctor_added_own_line_is_skipped(ctl, order_xlsx):
    ctl.import_xlsx(order_xlsx)
    text = ctl.generate("лигатурная").text + "\nДополнительно: беседа о гигиене."
    ctl.save_reference("лигатурная", text)
    assert ctl.workbook.references[-1].lines[-1] == ("", "Дополнительно: беседа о гигиене.")
    assert ctl.file_status == "✓ пунктов: 124 · эталоны 2/2"


def test_doctor_removed_row_is_mismatch(ctl, order_xlsx):
    ctl.import_xlsx(order_xlsx)
    text = ctl.generate("лигатурная").text.replace("осмотр.\n", "")
    ctl.save_reference("лигатурная", text)
    assert ctl.file_status == "⚠ эталоны 1/2 — подробнее"


# ---------- итоговое ревью 0.2b2 ----------
CHAIN = [("4p", "Эластическая цепочка без шага"), ("4p1", " на нижней челюсти."),
         ("4p2", "на верхней челюсти.")]


def test_two_variants_roundtrip(order_xlsx_with):
    wb = order_xlsx_with(*CHAIN)
    p = build_protocol(wb, "лигатурная, цепочка на обеих челюстях")
    assert [c for c, _ in p.lines if c.startswith("4p")] == ["4p", "4p1", "4p2"]
    assert lines_for_reference(wb, p, p.text) == p.lines


def test_normalize_drops_repeated_group(order_xlsx_with):
    wb = order_xlsx_with(*CHAIN)
    case = wb.references[0]
    case.lines += [("4p", "Эластическая цепочка без шага"), ("4p1", " на нижней челюсти."),
                   ("4p", "Эластическая цепочка без шага"), ("4p2", "на верхней челюсти.")]
    normalize_references(wb)
    assert [c for c, _ in case.lines if c.startswith("4p")] == ["4p", "4p1", "4p2"]


def test_normalize_full_sentence_variants(order_xlsx):
    """Врач переписал варианты целыми фразами (order-new-aktiv.xlsx), а эталон —
    из 0.1 со склеенной строкой: перевод находит вариант по словам после группы."""
    import openpyxl
    xl = openpyxl.load_workbook(order_xlsx)
    ws = xl.worksheets[0]
    new = {"2b1": "Тип системы: вестибулярная самолигируемая.",
           "2b2": "Тип системы: вестибулярная лигатурная.",
           "2c2": "Этап лечения:рабочий (коррекция окклюзии)."}
    for row in ws.iter_rows(min_col=1, max_col=2):
        if str(row[0].value).strip() in new:
            row[1].value = new[str(row[0].value).strip()]
    xl.save(order_xlsx)
    wb = load_workbook_file(order_xlsx)
    lines = wb.references[0].lines
    assert ("2b2", "Тип системы: вестибулярная лигатурная.") in lines
    assert ("2c2", "Этап лечения:рабочий (коррекция окклюзии).") in lines
    (r,) = check_references(wb)
    assert r.ok and r.own_lines == 0
