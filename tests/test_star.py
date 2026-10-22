"""0.5 beta 7: «*» в ID_CLASS — свойство по умолчанию сразу в отчёте (обязательное).
Пресет без звёздочек — как раньше; со звёздочками — в отчёт само идёт только отмеченное."""
import openpyxl
import pytest

from orthofill.engine import build_protocol
from orthofill.shellproto import shell_protocol
from orthofill.workbook import load_workbook_file


def _star(path, *codes):
    """Поставить «*» к кодам в колонке A (ID_CLASS)."""
    from orthofill.workbook import clean_code
    book = openpyxl.load_workbook(path)
    ws = book.worksheets[0]
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if cell.value is not None and clean_code(cell.value) in codes:
            cell.value = f'{cell.value}*'
    book.save(path)
    return load_workbook_file(path)


def codes(p):
    return [c for c, _ in p.lines]


def test_star_parsed(order_xlsx):
    wb = _star(order_xlsx, '4j', '3f1')
    assert '4j' in wb.by_id and wb.by_id['4j'].star and wb.by_id['3f1'].star
    assert not wb.by_id['2a'].star and wb.star_mode
    assert not wb.problems


def test_no_stars_unchanged(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    assert not wb.star_mode
    assert {'2a', '4j', '5d', '5d2'} <= set(codes(build_protocol(wb, '')))


def test_star_mode_only_starred(order_xlsx):
    wb = _star(order_xlsx, '2a', '4j', '3f1')
    empty = codes(build_protocol(wb, ''))
    assert {'2a', '4j', '3f', '3f1'} <= set(empty)
    assert '2d' not in empty and '5a' not in empty            # не отмечены — только если написано
    assert '2c' not in empty and '5d' not in empty            # умолчания-правила выключены


def test_star_mode_written_still_works(order_xlsx):
    wb = _star(order_xlsx, '4j')
    p = codes(build_protocol(wb, 'лигатурная, этап рабочий, перепел, явка 6 нед'))
    assert {'2b2', '2c2', '4f1', '4j'} <= set(p)


def test_group_star_default_variant(order_xlsx):
    wb = _star(order_xlsx, '5d2')
    p = build_protocol(wb, '')
    assert '5d2' in codes(p)
    assert '5d3' in codes(build_protocol(wb, 'явка через 6 недель')) or \
        '5d2' not in codes(build_protocol(wb, 'явка через 6 недель'))


def test_shell_mode_follows_stars(order_xlsx):
    wb = _star(order_xlsx, '4j')
    assert codes(shell_protocol(wb, '')) == codes(build_protocol(wb, ''))
    assert '2a' not in codes(shell_protocol(wb, ''))


def test_export_roundtrip_keeps_star(order_xlsx, tmp_path, store):
    from orthofill.xlsx_export import export_workbook
    wb = _star(order_xlsx, '4j', '3f1')
    pid = store.create(wb)
    loaded = store.load(pid)
    assert loaded.by_id['4j'].star and loaded.star_mode
    out = tmp_path / 'out.xlsx'
    export_workbook(loaded, out)
    again = load_workbook_file(out)
    assert again.by_id['4j'].star and again.by_id['3f1'].star and not again.by_id['2a'].star


def test_draft_set_star(order_xlsx, ctl):
    ctl.import_xlsx(order_xlsx)
    d = ctl.open_draft()
    d.set_prop('4j', 'star', True)
    ctl.save_draft()
    assert ctl.workbook.by_id['4j'].star and ctl.workbook.star_mode


def test_editor_checkbox(app):
    ed = app.open_editor('4j')
    app.root.update()
    assert ed.star_var is not None
    ed.star_var.set(True)
    ed._on_star()
    assert ed.draft.wb.by_id['4j'].star


def test_star_mode_explicit_always_kept(order_xlsx):
    """Явный выбор врача «всегда» в редакторе — тоже решение врача, не пропадает."""
    wb = _star(order_xlsx, '4j')
    wb.by_id['2d'].when = 'always'
    assert '2d' in codes(build_protocol(wb, ''))


def test_migrate_v4_db_gets_star(tmp_path):
    """База 0.5 b6 (версия 4) без колонки star — после открытия пресет грузится, звёздочек нет."""
    import sqlite3
    from orthofill.journal import Journal
    db = tmp_path / 'j.db'
    Journal(db)._db.close()
    con = sqlite3.connect(db)
    con.execute('ALTER TABLE preset_items DROP COLUMN star')
    con.execute('PRAGMA user_version = 4')
    con.commit(); con.close()
    j = Journal(db)
    cols = {r[1] for r in j._db.execute('PRAGMA table_info(preset_items)')}
    assert 'star' in cols
    assert j._db.execute('PRAGMA user_version').fetchone()[0] == 5
