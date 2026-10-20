import openpyxl
import pytest

from orthofill.workbook import WorkbookError, load_workbook_file


def make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def test_load_order(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    assert len(wb.items) == 124 and not wb.problems
    assert [s.id for s in wb.sections] == ['1', '2', '3', '4', '5']
    assert [c.id for c in wb.by_id['3'].children][:3] == ['3a', '3b', '3c']
    assert wb.by_id['4h1'].orphan and wb.by_id['4h1'].level == 2
    assert [c.id for c in wb.by_id['3y'].children] == ['3y1', '3y2', '3y3']
    assert len(wb.sha256) == 64


def test_references(order_xlsx):
    (case,) = load_workbook_file(order_xlsx).references
    assert case.case_id == '1' and case.raw.startswith('активация, лигатурная')
    assert case.lines[0] == ('1', 'Жалобы: активация')
    assert ('2b', 'Тип системы вестибулярная лигатурная.') in case.lines


def test_load_messy_codes(tmp_path):
    p = make_xlsx(tmp_path / 'm.xlsx', [
        ('ID_CLASS', 'CLASS_LIST'),
        (None, None),
        (1.0, 'Жалобы:'),
        ('1a', 'нет<br>жалоб'),
        (4, 'Лечение:'),
        (' 4B ', 'Смена дуги'),
        ('4b1', 'на верхней 0.014" NiTi'),
    ])
    wb = load_workbook_file(p)
    assert not wb.problems
    assert {'1', '1a', '4', '4b', '4b1'} <= set(wb.by_id)
    assert wb.by_id['1a'].text == 'нет жалоб'
    assert [c.id for c in wb.by_id['4b'].children] == ['4b1']


def test_problems(tmp_path):
    p = make_xlsx(tmp_path / 'p.xlsx', [
        ('ID_CLASS', 'CLASS_LIST'),
        ('1', 'Жалобы:'),
        ('1a', 'нет'),
        ('1a', 'дубль'),
        ('1b', None),
        ('zz', 'плохой код'),
    ])
    probs = load_workbook_file(p).problems
    assert any('строка 4' in x and '1a' in x for x in probs)
    assert any('строка 5' in x and '1b' in x for x in probs)
    assert any('строка 6' in x and 'zz' in x for x in probs)


def test_no_header(tmp_path):
    p = make_xlsx(tmp_path / 'n.xlsx', [('1', 'Жалобы')])
    with pytest.raises(WorkbookError):
        load_workbook_file(p)


def test_missing_file(tmp_path):
    with pytest.raises(WorkbookError):
        load_workbook_file(tmp_path / 'нет.xlsx')
