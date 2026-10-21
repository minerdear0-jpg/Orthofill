import openpyxl

from orthofill.engine import ReferenceCase, build_protocol, check_references, rule_kind, rule_warnings
from orthofill.workbook import load_workbook_file

REF_RAW = ("активация, лигатурная,этап рабочий, класс по молярам 1 и клыкам, замена дуги "
           "на верху 18 25 нити. ношение эластиков перепел, явка через 4 недели. ")


def ids(protocol):
    return [i for i, _ in protocol.lines]


def test_reference_order_xlsx(order_xlsx):
    assert [r.ok for r in check_references(load_workbook_file(order_xlsx))] == [True]


def test_reference_compares_text(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    case = wb.references[0]
    lines = list(case.lines)
    lines[1] = (lines[1][0], 'другой текст')
    wb.references = [ReferenceCase(case.case_id, case.raw, lines)]
    (res,) = check_references(wb)
    assert not res.ok and ('1a', 'другой текст') in res.missing


def test_text_is_plain(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "лигатурная")
    assert p.text.splitlines()[0] == "Жалобы: активация" and "\t" not in p.text
    assert "1a" not in p.text


def test_defaults_reported(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "лигатурная")
    assert {"этап рабочий", "явка 4 нед", "жалоб нет", "лигатуры на обеих челюстях"} <= set(p.defaults)


def test_no_defaults_when_written(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), REF_RAW)
    assert "этап рабочий" not in p.defaults and "явка 4 нед" not in p.defaults


def test_rows_merge_and_lines_per_item(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "самолигирующая, детализация, явка 6 нед")
    assert ("2b", "Тип системы") in p.lines and ("2b1", "вестибулярная самолигируемая.") in p.lines
    assert "Тип системы вестибулярная самолигируемая." in p.text.splitlines()
    assert "Явка на следующую активацию через 6 недель." in p.text.splitlines()
    assert "4h2" in ids(p) and "4h1" not in ids(p)
    assert "\n\n" not in p.text


def test_section_rows(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "лигатурная")
    assert [r.text for r in p.rows if r.section] == [
        "Жалобы: активация", "Анамнез: активация", "Объективно: активация", "Лечение:", "Рекомендации:"]


def test_check_marks(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx),
                       "лигатурная, по молярам 1, замена дуги 18х25 нити, явка 5 нед")
    marked = {c for r in p.rows if r.check for c in r.codes}
    assert {"3g", "3h", "4b", "4c", "5d"} <= marked
    assert not any(r.check for r in p.rows if "3f" in r.codes)      # по умолчанию — без метки


def test_no_marks_when_explicit(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx),
                       "лигатурная, замена дуги верх 18х25 нити, явка 6 нед")
    assert not any(r.check for r in p.rows)


def test_new_elastic_option(order_xlsx_with):
    p = build_protocol(order_xlsx_with(("4f7", "Слон")), "лигатурная, резинки слон")
    assert ("4f7", "Слон") in p.lines


def test_new_visit_option(order_xlsx_with):
    p = build_protocol(order_xlsx_with(("5d5", "3 недели.")), "лигатурная, явка через 3 недели")
    assert ("5d5", "3 недели.") in p.lines and not p.warnings


def test_new_complaint(order_xlsx_with):
    p = build_protocol(order_xlsx_with(("1e", "на боль в области клыка.")),
                       "активация, жалобы на боль клык")
    assert p.lines[1][0] == "1e"


def test_known_complaint_not_confused_by_wire_change(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "замена дуги верх 16 нити")
    assert p.lines[1][0] == "1a"
    p = build_protocol(load_workbook_file(order_xlsx), "отклеился брекет")
    assert p.lines[1][0] == "1b"


def test_renumbered_group(order_xlsx):
    xl = openpyxl.load_workbook(order_xlsx)
    ws = xl.worksheets[0]
    for (c,) in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(c.value, str) and c.value.startswith('4f'):
            c.value = '4m' + c.value[2:]
    xl.save(order_xlsx)
    wb = load_workbook_file(order_xlsx)
    assert rule_kind(wb.by_id['4m']) == 'elastics'
    assert ("4m1", "Перепел") in build_protocol(wb, "эластики перепел").lines


def test_unknown_group_matched_by_words(order_xlsx_with):
    wb = order_xlsx_with(("4z", "Ретейнер"), ("4z1", "несъёмный установлен."),
                         ("4z2", "съёмный выдан."))
    assert ("4z2", "съёмный выдан.") in build_protocol(wb, "ретейнер съемный выдан").lines
    assert "4z" not in ids(build_protocol(wb, "лигатурная"))


def test_unknown_elastic_warns(order_xlsx):
    assert "эластики упомянуты, но вариант не найден в классификаторе" in \
        build_protocol(load_workbook_file(order_xlsx), "резинки бегемот").warnings


def test_wire_warnings(order_xlsx):
    w = build_protocol(load_workbook_file(order_xlsx), "замена дуги верх 14x25 ss, низ нити").warnings
    assert 'дуга на верхней: 0.014×0.025" SS нет в классификаторе' in w
    assert "дуга на нижней: размер не распознан — выберите вручную" in w


def test_two_wires(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "замена дуги верх 16х25 СuNiTi, низ 0,014 нити")
    assert {"4b10", "4c1"} <= set(ids(p))


def test_visit_nearest_warns(order_xlsx):
    p = build_protocol(load_workbook_file(order_xlsx), "явка через 5 нед")
    assert any(w.startswith("явка через 5 нед: такого варианта нет") for w in p.warnings)


def test_class_three(order_xlsx):
    assert {"3g3", "3h3"} <= set(ids(build_protocol(load_workbook_file(order_xlsx), "по молярам 3")))


def test_rule_warnings(order_xlsx):
    wb = load_workbook_file(order_xlsx)
    assert rule_warnings(wb) == []
    wb.by_id['4f'].text = 'Резинки'
    assert "в файле нет группы для правила «эластики»" in rule_warnings(wb)
