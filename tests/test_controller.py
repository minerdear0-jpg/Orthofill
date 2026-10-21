import shutil
import subprocess
import sys

import pytest

from orthofill.controller import Controller, layout_mode
from orthofill.journal import Journal
from orthofill.workbook import FileLockedError

from conftest import ORDER, ROOT

REF_RAW = ("активация, лигатурная,этап рабочий, класс по молярам 1 и клыкам, замена дуги "
           "на верху 18 25 нити. ношение эластиков перепел, явка через 4 недели. ")


@pytest.fixture
def clip():
    return []


@pytest.fixture
def ctl(tmp_path, clip):
    return Controller(Journal(tmp_path / "j.db"), clip.append)


def test_layout_mode():
    assert layout_mode(899) == "narrow" and layout_mode(900) == "wide"


def test_no_file(ctl):
    assert ctl.file_status == "Файл не открыт"
    assert ctl.generate("лигатурная").status == "Сначала откройте файл классификатора"


def test_file_status_ok(ctl, order_xlsx):
    ctl.open_file(order_xlsx)
    assert ctl.file_status == "✓ пунктов: 124 · эталоны 1/1"
    assert ctl.journal.get_setting("last_file") == str(order_xlsx)


def test_file_status_mismatch(ctl, order_xlsx):
    from orthofill.workbook import append_reference
    append_reference(order_xlsx, "лигатурная", [("1", "Жалобы: активация")])
    ctl.open_file(order_xlsx)
    assert ctl.file_status == "⚠ эталоны 1/2 — подробнее"
    assert any("эталон 2" in d for d in ctl.details)


def test_file_with_problems_blocks_generate(ctl, tmp_path):
    import openpyxl
    p = tmp_path / "bad.xlsx"
    xl = openpyxl.Workbook()
    for r in [("ID_CLASS", "CLASS_LIST"), ("1", "Жалобы:"), ("1", "дубль")]:
        xl.active.append(r)
    xl.save(p)
    ctl.open_file(p)
    assert ctl.file_status == "✗ ошибки в файле: 1 — подробнее"
    assert ctl.generate("x").status == "Файл содержит ошибки — откройте «подробнее»"


def test_generate_copies_and_logs(ctl, clip, order_xlsx):
    ctl.open_file(order_xlsx)
    r = ctl.generate(REF_RAW)
    assert r.copied and r.status == "Скопировано в буфер" and clip == [r.text]
    assert r.text.splitlines()[0] == "Жалобы: активация"
    assert "этап рабочий" not in r.defaults_line and "явка" not in r.defaults_line
    assert ctl.last_entry_id is not None


def test_generate_defaults_line(ctl, order_xlsx):
    ctl.open_file(order_xlsx)
    assert ctl.generate("лигатурная").defaults_line.startswith("по умолчанию: жалоб нет, этап рабочий")


def test_generate_empty_raw(ctl, clip, order_xlsx):
    ctl.open_file(order_xlsx)
    r = ctl.generate("   \n")
    assert r.status == "Введите краткую запись" and not r.copied and clip == []


def test_startup_missing_last_file(ctl, tmp_path):
    ctl.journal.set_setting("last_file", str(tmp_path / "пропал.xlsx"))
    ctl.startup()
    assert ctl.workbook is None
    assert ctl.file_status == "Файл не найден — откройте файл классификатора"


def test_startup_opens_last_file(ctl, order_xlsx):
    ctl.journal.set_setting("last_file", str(order_xlsx))
    ctl.startup()
    assert ctl.workbook is not None


def test_regenerate_needs_confirm_after_edit(ctl, order_xlsx):
    ctl.open_file(order_xlsx)
    assert not ctl.needs_confirm_regenerate("")
    text = ctl.generate("лигатурная").text
    assert not ctl.needs_confirm_regenerate(text)
    assert ctl.needs_confirm_regenerate(text + "\nправка")
    ctl.copy(text + "\nправка")
    assert not ctl.needs_confirm_regenerate(text + "\nправка")


def test_copy_logs(ctl, clip, order_xlsx):
    ctl.open_file(order_xlsx)
    ctl.generate("лигатурная")
    assert ctl.copy("правленый") == "Скопировано в буфер" and clip[-1] == "правленый"


def test_save_reference(ctl, order_xlsx):
    ctl.open_file(order_xlsx)
    text = ctl.generate("лигатурная, перепел").text
    assert ctl.save_reference("лигатурная, перепел", text) == "Эталон 2 сохранён в order.xlsx"
    assert ctl.file_status == "✓ пунктов: 124 · эталоны 2/2"


def test_save_reference_locked(ctl, order_xlsx, monkeypatch):
    import orthofill.controller as cmod
    ctl.open_file(order_xlsx)
    text = ctl.generate("лигатурная").text

    def locked(*a, **k):
        raise FileLockedError()
    monkeypatch.setattr(cmod, "append_reference", locked)
    assert ctl.save_reference("лигатурная", text) == "Закройте файл в Excel и повторите"


def test_cli_test_mode():
    r = subprocess.run([sys.executable, "-m", "orthofill", "--test", str(ORDER)],
                       cwd=ROOT, capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "[OK]" in r.stdout


def test_cli_raw_mode():
    r = subprocess.run([sys.executable, "-m", "orthofill", "--xlsx", str(ORDER), "--raw",
                        "замена дуги верх 14x25 ss"], cwd=ROOT, capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0 and r.stdout.startswith("Жалобы: активация")
    assert "нет в классификаторе" in r.stderr


@pytest.mark.skipif(sys.platform != "win32" and not (
        __import__("os").environ.get("DISPLAY") or __import__("os").environ.get("WAYLAND_DISPLAY")),
    reason="нужен дисплей")
def test_cli_selftest_ui():
    r = subprocess.run([sys.executable, "-m", "orthofill", "--selftest-ui"],
                       cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert r.returncode == 0, r.stderr


# ---------- ревью: I7, I8 ----------
def test_save_reference_refuses_changed_note(ctl, order_xlsx):
    ctl.open_file(order_xlsx)
    text = ctl.generate("лигатурная, перепел").text
    assert ctl.save_reference("лигатурная, лиса", text) == \
        "Запись изменена — сформируйте протокол заново"
    from orthofill.workbook import load_workbook_file
    assert len(load_workbook_file(order_xlsx).references) == 1


def test_save_reference_unexpected_error_is_shown(ctl, order_xlsx, monkeypatch):
    import orthofill.controller as cmod
    ctl.open_file(order_xlsx)
    text = ctl.generate("лигатурная").text

    def boom(*a, **k):
        raise OSError("диск переполнен")
    monkeypatch.setattr(cmod, "append_reference", boom)
    assert ctl.save_reference("лигатурная", text) == "✗ Ошибка: диск переполнен"


def test_open_file_unexpected_error_is_shown(ctl, order_xlsx, monkeypatch):
    import orthofill.controller as cmod

    def boom(*a, **k):
        raise ValueError("битый файл")
    monkeypatch.setattr(cmod, "load_workbook_file", boom)
    ctl.open_file(order_xlsx)
    assert ctl.workbook is None and ctl.file_status == "✗ Ошибка: битый файл"


def test_generate_unexpected_error_is_shown(ctl, clip, order_xlsx, monkeypatch):
    import orthofill.controller as cmod
    ctl.open_file(order_xlsx)

    def boom(*a, **k):
        raise RuntimeError("сбой")
    monkeypatch.setattr(cmod, "build_protocol", boom)
    r = ctl.generate("лигатурная")
    assert r.status == "✗ Ошибка: сбой" and not r.copied and clip == []


def test_save_reference_file_renamed(ctl, order_xlsx):
    """Диагностика 21.10: файл переименовали, пока программа была открыта."""
    ctl.open_file(order_xlsx)
    text = ctl.generate("лигатурная").text
    order_xlsx.rename(order_xlsx.with_name("order-new-aktiv.xlsx"))
    msg = ctl.save_reference("лигатурная", text)
    assert "не найден" in msg and "Открыть файл" in msg


def test_cli_check_bad_file_is_friendly(tmp_path):
    import openpyxl
    import subprocess
    import sys
    p = tmp_path / "kons.xlsx"
    openpyxl.Workbook().save(p)
    r = subprocess.run([sys.executable, "-m", "orthofill", "--check", str(p)],
                       capture_output=True, encoding="utf-8", cwd=ROOT)
    assert r.returncode == 2 and "Traceback" not in r.stderr
    assert "ID_CLASS" in r.stdout + r.stderr
