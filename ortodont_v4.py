# -*- coding: utf-8 -*-
"""
ortodont.py — маппинг raw-записи врача на ID_CLASS классификатора.

Режимы:
  python ortodont.py                              # GUI
  python ortodont.py --check order.xlsx           # только проверка (exit 0/1)
  python ortodont.py --test order.xlsx            # сверка с эталонами C-F
  python ortodont.py --xlsx order.xlsx --raw "..." [--out result.txt]
  python ortodont.py --xlsx order.xlsx --raw-file raw.txt [--format tsv|pipe|doc]

Формат order.xlsx (лист 1):
  A = ID_CLASS  (1, 1a, 2b1, 4b10, ...)
  B = CLASS_LIST (текст классификатора)
"""
import sys, re, json, argparse
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Установите: pip install openpyxl")

# ============================================================
# КОНФИГ
# ============================================================
DEFAULT_CONFIG = {
    # Буалярные ID: выводятся всегда при активной секции
    "fixed_ids": ["2a", "2d", "3a", "3b", "3c", "3d",
                  "4a", "4j", "5a", "5b", "5c"],
    # ID, где текст родителя склеивается с текстом выбранного ребёнка
    "combine_parents": ["2b", "2c"],
    # Секции, где из детей выбирается один (по смыслу — жалобы)
    "choose_one_sections": {"1": "complaint"},
    # Порог нечёткого совпадения для «сиротских» и новых листовых позиций
    "threshold": 0.25,
}

def load_config(path=None):
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    p = Path(path) if path else Path(__file__).with_name("ortodont_config.json")
    if p.exists():
        try:
            data = json.loads(p.read_text("utf-8"))
            for k, v in data.items():
                if isinstance(v, dict) and isinstance(cfg.get(k), dict):
                    cfg[k].update(v)
                else:
                    cfg[k] = v
        except Exception as e:
            print(f"config: {e}", file=sys.stderr)
    return cfg

# ============================================================
# ТЕКСТ
# ============================================================
# Слово, набранное в двух раскладках («СuNiTi», «NiТi», «нижнeй»), приводим
# к той раскладке, букв которой в нём больше.
_TO_LAT = str.maketrans('АВСЕНКМОРТХУИаеорсхутнкмиі',
                        'ABCENKMOPTXUIaeopcxutnkmii')
_TO_CYR = str.maketrans('ABCEHKMOPTXYaeopcxykmth',
                        'АВСЕНКМОРТХУаеорсхукмтн')

def _fix_mixed_word(m):
    w = m.group(0)
    cyr = len(re.findall(r'[а-яёА-ЯЁіІ]', w))
    lat = len(re.findall(r'[a-zA-Z]', w))
    if not cyr or not lat:
        return w
    return w.translate(_TO_LAT) if lat >= cyr else w.translate(_TO_CYR)

def norm(t):
    if t is None:
        return ''
    s = re.sub(r'[A-Za-zА-Яа-яЁёІі]+', _fix_mixed_word, str(t))
    s = s.lower().replace('ё', 'е').replace('×', 'x')
    s = s.replace('–', '-').replace('—', '-')
    s = re.sub(r'<br\s*/?>', ' ', s, flags=re.I)
    return re.sub(r'\s+', ' ', s).strip()

def strip_html(t):
    if not t:
        return ''
    return re.sub(r'<br\s*/?>', ' ', str(t), flags=re.I).strip()

def natural_key(s):
    return [(0, int(x)) if x.isdigit() else (1, x)
            for x in re.split(r'(\d+)', s) if x]

# ============================================================
# ПАРСЕРЫ RAW
# ============================================================
def parse_system_type(n):
    if 'самолигир' in n:
        return 'self_ligating'
    if 'лигатур' in n:
        return 'ligature'
    return None

def parse_stage(n):
    if 'детализ' in n:
        return 'detail'
    if 'начальн' in n or 'нивелир' in n:
        return 'initial'
    return 'working'

def clause(n, pattern, stops='.;'):
    """Фрагмент raw от совпадения pattern до ближайшего знака из stops."""
    m = re.search(pattern, n)
    if not m:
        return ''
    # знак внутри числа («0.014», «0,018») — не конец фрагмента
    stop = re.compile('[' + re.escape(stops) + r'](?!\d)')
    s = stop.search(n, m.end())
    return n[m.start():s.start() if s else len(n)]

# Упоминание челюсти: верх/верхняя/сверху/в/ч/вч, низ/нижняя/снизу/н/ч/нч, обе/обеих/обоих
ARCH_RE = re.compile(
    r'(?P<both>\bоб(?:е|еих|оих)\b)'
    r'|(?P<upper>\bс?верх\w*|\bв\s*/\s*ч\b|\bвч\b)'
    r'|(?P<lower>\bс?ни[зж]\w*|\bн\s*/\s*ч\b|\bнч\b)')

def arch_marks(text):
    """[(позиция начала, позиция конца, 'upper'|'lower'|'both'), ...]"""
    return [(m.start(), m.end(), m.lastgroup) for m in ARCH_RE.finditer(text)]

def arches_in(fragment):
    found = {a for _, _, a in arch_marks(fragment)}
    if 'both' in found or {'upper', 'lower'} <= found:
        return 'both'
    return found.pop() if found else None

_CLASS_NUM = r'([1-3]|i{1,3})(?![\w])'

def _class_val(v):
    return {'i': 1, 'ii': 2, 'iii': 3}.get(v) or int(v)

def parse_classes(n):
    """Класс моляров и клыков. Число относится к слову, рядом с которым стоит:
    «по молярам 1», «моляры - II», «1 класс по клыкам»."""
    result = {'molar': None, 'canine': None}
    keys = {'моляр': 'molar', 'клык': 'canine'}
    # слово → число («по молярам 1 и клыкам»: 1 — только моляры)
    for m in re.finditer(r'(моляр|клык)\w*\s*[-:]?\s*(?:по\s+)?' + _CLASS_NUM, n):
        k = keys[m.group(1)]
        if result[k] is None:
            result[k] = _class_val(m.group(2))
    # число → слово («1 класс по молярам»)
    for m in re.finditer(r'(?<![\w.,])' + _CLASS_NUM +
                         r'\s*(?:класс\w*\s*)?(?:по\s+)?(моляр|клык)', n):
        k = keys[m.group(2)]
        if result[k] is None:
            result[k] = _class_val(m.group(1))
    return result

# ---------- ДУГИ ----------
# Число размера: 0.016 / 0,016 / .016 / 016 / 16  → «16»
_NUM = r'(?:0?[.,])?0?(\d{2})(?!\d)'
# Разделитель сечения: x (лат.), х (кир.), *, -, /, «на», точка, запятая, пробел
_SIZE2 = re.compile(r'(?<![\d.,])' + _NUM +
                    r'\s*(?:x|х|\*|-|/|на|[.,]|\s)\s*' + _NUM)
_SIZE1 = re.compile(r'(?<![\d.,])' + _NUM)

# Порядок важен: CuNiTi проверяем раньше NiTi («кунити» содержит «нити»)
_MATERIALS = [
    ('CuNiTi', r'c[uy]\s*-?\s*ni\s*-?\s*ti|ку\s*-?\s*ни\s*-?\s*ти|медн\w*|купрум'),
    ('TMA',    r'\btma\b|\bтма\b|бета\s*-?\s*титан|\bbeta\b|β|молибд'),
    ('SS',     r'\bss\b|\bсс\b|стал\w*|\bsteel\b|нерж\w*'),
    ('NiTi',   r'(?<![a-z])ni\s*-?\s*ti|(?<![а-я])ни\s*-?\s*ти|никел\w*'),
]

def detect_material(text):
    for name, rx in _MATERIALS:
        if re.search(rx, text):
            return name
    return None

def detect_size(text):
    """Первое правдоподобное сечение: круглое 12–21 или прямоугольное
    (12–21)×(22–28): глубина у ортодонтических дуг .022/.025/.028."""
    for m in _SIZE2.finditer(text):
        a, b = int(m.group(1)), int(m.group(2))
        if 12 <= a <= 21 and 22 <= b <= 28:
            return (m.group(1), m.group(2))
    for m in _SIZE1.finditer(text):
        if 12 <= int(m.group(1)) <= 21:
            return (m.group(1), None)
    return None

# Пункт записи про другое действие — не часть фразы о дуге
_WIRE_STOP = re.compile(
    r'лигатур|подвяз|эластик|явк|изгиб|юстир|гигиен|слизист|эрози|брекет|'
    r'отклеи|выход|перелом|динамик|класс|моляр|клык|этап|зуб|переклей|'
    r'пружин|недел')
_NEGATION = re.compile(r'без\s+(?:замен|смен|изменен)|не\s+мен|оставл|прежн')

def _wire_items(n):
    """Текст про дуги: непрерывная цепочка пунктов (через , ; . — но не
    внутри числа), где есть «дуг», размер, материал или челюсть.
    Пункт про лигатуры, эластики, явку и т.п. цепочку обрывает."""
    items = [s.strip() for s in re.split(r'[;\n]|[.,](?!\d)', n)]
    runs, cur, has_wire = [], [], False
    for it in items + ['']:
        wireish = bool(it) and not _WIRE_STOP.search(it) and (
            'дуг' in it or detect_size(it) or detect_material(it)
            or arch_marks(it))
        if wireish:
            cur.append(it)
            has_wire |= 'дуг' in it
        else:
            if cur and has_wire:
                runs.append(', '.join(cur))
            cur, has_wire = [], False
    return runs

def _spec(text):
    if _NEGATION.search(text):
        return 'skip'
    size, mat = detect_size(text), detect_material(text)
    return (size, mat) if (size or mat) else None

def parse_wire_from_raw(n):
    """{'upper'|'lower': ((ширина, глубина|None) | None, материал | None)}

    Текст режем по упоминаниям челюстей. Пример: «замена дуги верх 16х25 нити,
    низ 14 нити» → [префикс «замена дуги»] [верх: 16х25 нити] [низ: 14 нити].
    Если размер стоит ДО челюсти («16 нити верх, 14 низ»), сдвигаем на один."""
    result = {}
    for region in _wire_items(n):
        marks = arch_marks(region)
        bounds = [m[0] for m in marks] + [len(region)]
        prefix = _spec(region[:bounds[0]])
        after = [_spec(region[m[1]:bounds[i + 1]]) for i, m in enumerate(marks)]

        if not marks:
            if prefix and prefix != 'skip':
                result['upper'] = result['lower'] = prefix
            continue

        # «16 нити верх, 14 низ»: размер перед челюстью, после последней — пусто
        if prefix and prefix != 'skip' and after[-1] is None:
            specs = [prefix] + after[:-1]
        else:
            specs = after
        # Челюсть без своего размера берёт соседний: «верх и низ 16х25 нити»
        for i, s in enumerate(specs):
            if s is None:
                nxt = next((x for x in specs[i + 1:] if x not in (None, 'skip')), None)
                prv = next((x for x in reversed(specs[:i]) if x not in (None, 'skip')),
                           prefix if prefix != 'skip' else None)
                specs[i] = nxt or prv
        for (_, _, arch), s in zip(marks, specs):
            if s in (None, 'skip'):
                continue
            for a in (('upper', 'lower') if arch == 'both' else (arch,)):
                result[a] = s

    # Материал назван один раз на всё: «верх 16x25, низ 14, нити»
    mats = {m for _, m in result.values() if m}
    if len(mats) == 1:
        mat = mats.pop()
        result = {a: (s, m or mat) for a, (s, m) in result.items()}
    return result

def parse_wire_from_label(label):
    n = norm(label)
    arch = 'upper' if 'верхн' in n else 'lower' if 'нижн' in n else None
    return arch, detect_size(n), detect_material(n)

def parse_elastics(n):
    for animal, key in [
        ('перепел', 'перепел'), ('перепёл', 'перепел'),
        ('лиса', 'лиса'), ('лису', 'лиса'),
        ('медвед', 'медведь'),
        ('антилоп', 'антилопа'),
        ('кролик', 'кролик'),
        ('зебра', 'зебра'), ('зебру', 'зебра'),
    ]:
        if animal in n:
            return key
    return None

def parse_weeks(n):
    m = re.search(r'через\s+(\d+)\s*(?:недел|нед)', n)
    if m:
        return int(m.group(1))
    m = re.search(r'(\d+)\s*(?:недел|нед)', n)
    if m:
        return int(m.group(1))
    return 4

def parse_hygiene(n):
    return 'unsatisfactory' if 'гигиен' in n and 'неудовлетвор' in n else 'satisfactory'

def parse_mucosa(n):
    return 'erosion' if 'эрози' in n else 'normal'

def parse_bracket_state(n):
    return 'unsatisfactory' if 'брекет' in n and 'неудовлетвор' in n else 'satisfactory'

def parse_dynamics(n):
    if re.search(r'без\s+(?:\w+\s+)?(?:динамик|существен)|нет\s+динамик', n):
        return 'no_dynamics'
    return 'plan'

def parse_ligature_change(n, st):
    if re.search(r'(смен|замен)\w*\s+[^.;]*лигатур|лигатур\w*\s+(смен|замен)', n):
        return True
    return st == 'ligature'

def parse_ligature_arches(n):
    """Лигатуры меняют на обеих челюстях, если явно не сказано иначе.
    Челюсть ищем только во фрагменте про лигатуры, а не во всём тексте
    (иначе «дуга на верху» уводит лигатуры на верхнюю)."""
    frag = clause(n, r'лигатур(?!н)', stops='.;,')
    return arches_in(frag) or 'both'

def parse_binding_arches(n):
    return arches_in(clause(n, r'подвяз', stops='.;,'))

def parse_wire_adjustments(n):
    out = []
    if re.search(r'изгиб\w*\s+перв|перв\w*\s+порядк', n):
        out.append(1)
    if re.search(r'изгиб\w*\s+втор|втор\w*\s+порядк', n):
        out.append(2)
    if re.search(r'изгиб\w*\s+трет|трет\w*\s+порядк', n):
        out.append(3)
    return out

def extract_facts(raw):
    n = norm(raw)
    st = parse_system_type(n)
    return {
        'raw_norm': n,
        'system_type': st,
        'stage': parse_stage(n),
        'classes': parse_classes(n),
        'wires': parse_wire_from_raw(n),
        'elastics': parse_elastics(n),
        'weeks': parse_weeks(n),
        'hygiene': parse_hygiene(n),
        'mucosa': parse_mucosa(n),
        'bracket_state': parse_bracket_state(n),
        'dynamics': parse_dynamics(n),
        'ligature_change': parse_ligature_change(n, st),
        'ligature_arches': parse_ligature_arches(n),
        'binding_arches': parse_binding_arches(n),
        'wire_adjustments': parse_wire_adjustments(n),
        'has_binding': 'подвяз' in n,
    }

# ============================================================
# КЛАССИФИКАТОР
# ============================================================
def load_classifier(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    items = []
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        code, text = row[0], row[1]
        if code is None or text is None:
            continue
        cs = str(code).strip()
        if cs.endswith('.0') and cs[:-2].isdigit():
            cs = cs[:-2]
        if not re.fullmatch(r'\d|\d[a-zа-я]|\d[a-zа-я]\d+', cs):
            continue
        txt = strip_html(text)
        if not txt:
            continue
        items.append({'id': cs, 'text': txt})
    return items

def build_hierarchy(items):
    by_id = {it['id']: it for it in items}
    for it in items:
        cs = it['id']
        if re.fullmatch(r'\d', cs):
            it['level'], it['parent'] = 1, None
        elif re.fullmatch(r'\d[a-zа-я]', cs):
            it['level'], it['parent'] = 2, cs[0]
        elif re.fullmatch(r'\d[a-zа-я]\d+', cs):
            it['level'] = 3
            it['parent'] = cs[:re.search(r'\d+$', cs).start()]
        else:
            it['level'], it['parent'] = 0, None
        it['original_level'] = it['level']
        it['children'] = []

    # Осиротевшие level-3 → повышаем до level-2 (самодостаточные)
    for it in items:
        if it['level'] == 3 and it['parent'] not in by_id:
            it['level'] = 2
            it['parent'] = it['id'][0]

    for it in items:
        if it['parent'] and it['parent'] in by_id:
            by_id[it['parent']]['children'].append(it)
    # порядок детей = порядок строк в xlsx (там 3y идёт раньше 3j)
    return by_id

# ============================================================
# МАТЧИНГ
# ============================================================
def match_class_num(text, num):
    # norm() уже в нижнем регистре: римские цифры тут «i», «ii», «iii»
    m = re.search(r'\b(i{1,3}|iv|[1-4])\s*класс', norm(text))
    if not m:
        return False
    v = m.group(1)
    if v[0] in 'iv':
        return {'i': 1, 'ii': 2, 'iii': 3, 'iv': 4}[v] == num
    return int(v) == num

def match_order(text, order):
    return {1: 'перв', 2: 'втор', 3: 'трет'}[order] in norm(text)

def simple_score(text, raw_norm):
    kws = [w for w in re.findall(r'[а-яa-z]{3,}|\d+', norm(text))]
    if not kws:
        return 0
    return sum(1 for k in kws if k in raw_norm) / len(kws)

def score_orphan(text, facts):
    n = norm(text)
    st = facts['system_type'] or 'ligature'
    if 'фиксаци' in n and 'лигатур' in n:
        return 1.0 if st != 'self_ligating' else 0.0
    if 'клипс' in n:
        return 1.0 if st == 'self_ligating' else 0.0
    return simple_score(text, facts['raw_norm'])

def select_complaint(children, facts):
    n = facts['raw_norm']
    for c in children:
        ct = norm(c['text'])
        if 'отклеива' in ct and re.search(r'отклеива|отвалил', n):
            return c
        if 'выход' in ct and 'дуг' in ct and re.search(r'выход\s+дуги', n):
            return c
        if 'перелом' in ct and re.search(r'перелом\s+дуг|сломал\w*\s+дуг', n):
            return c
    return children[0] if children else None

def select_child(parent, facts):
    pid = parent['id']
    children = parent['children']
    if not children:
        return None

    if pid == '2b':
        st = facts['system_type']
        if st is None:
            return None
        for c in children:
            ct = norm(c['text'])
            if st == 'self_ligating' and 'самолигир' in ct:
                return c
            if st == 'ligature' and 'лигатур' in ct and 'самолигир' not in ct:
                return c
        return None

    if pid == '2c':
        stg = facts['stage']
        for c in children:
            ct = norm(c['text'])
            if stg == 'detail' and 'детализ' in ct:
                return c
            if stg == 'initial' and ('начальн' in ct or 'нивелир' in ct):
                return c
            if stg == 'working' and 'рабоч' in ct:
                return c
        return None

    if pid == '3e':
        if facts['mucosa'] == 'erosion':
            for c in children:
                if 'эрози' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'без видимых' in ct or 'норм' in ct:
                return c
        return children[0]

    if pid == '3f':
        if facts['hygiene'] == 'unsatisfactory':
            for c in children:
                if 'неудовлетвор' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'удовлетвор' in ct and 'неудовлетвор' not in ct:
                return c
        return children[0]

    if pid in ('3g', '3h', '3y', '3j'):
        tooth = 'моляр' if pid in ('3g', '3h') else 'клык'
        cls = facts['classes']['molar' if tooth == 'моляр' else 'canine']
        if cls is None:
            return None
        side = 'справа' if pid in ('3g', '3y') else 'слева'
        for c in children:
            ct = norm(c['text'])
            if side in ct and tooth[:5] in ct and match_class_num(ct, cls):
                return c
        return None

    if pid == '3k':
        if facts['bracket_state'] == 'unsatisfactory':
            for c in children:
                if 'неудовлетвор' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'удовлетвор' in ct and 'неудовлетвор' not in ct:
                return c
        return children[0]

    if pid in ('4b', '4c'):
        arch = 'upper' if pid == '4b' else 'lower'
        wire_info = facts['wires'].get(arch)
        if wire_info is None:
            return None
        size, material = wire_info
        jaw = 'верхней' if arch == 'upper' else 'нижней'
        warn = facts.setdefault('warnings', []).append
        if size is None:
            warn(f"дуга на {jaw}: размер не распознан — выберите вручную")
            return None
        for c in children:
            la, ls, lm = parse_wire_from_label(c['text'])
            if la != arch or ls != size:
                continue
            if material and lm and material != lm:
                continue
            return c
        dims = '×'.join(f'0.0{x}' for x in size if x)
        warn(f"дуга на {jaw}: {dims}\" {material or ''} нет в классификаторе")
        return None

    if pid == '4d':
        if not facts['ligature_change']:
            return None
        arches = facts['ligature_arches']
        for c in children:
            ct = norm(c['text'])
            if arches == 'both' and 'обе' in ct:
                return c
            if arches == 'upper' and 'верхн' in ct:
                return c
            if arches == 'lower' and 'нижн' in ct:
                return c
        return None

    if pid == '4e':
        if not facts['has_binding']:
            return None
        arches = facts['binding_arches']
        for c in children:
            ct = norm(c['text'])
            if arches == 'upper' and 'верхн' in ct:
                return c
            if arches == 'lower' and 'нижн' in ct:
                return c
        return None

    if pid == '4f':
        el = facts['elastics']
        if el is None:
            return None
        for c in children:
            if el in norm(c['text']):
                return c
        return None

    if pid == '4g':
        if not facts['wire_adjustments']:
            return None
        for order in facts['wire_adjustments']:
            for c in children:
                if match_order(c['text'], order):
                    return c
        return None

    if pid == '4k':
        want = 'без' if facts['dynamics'] == 'no_dynamics' else 'соответств'
        for c in children:
            if want in norm(c['text']):
                return c
        return children[0]

    if pid == '5d':
        w = facts['weeks']
        for c in children:
            if str(w) in norm(c['text']):
                return c
        best, bd = None, 999
        for c in children:
            cm = re.search(r'(\d+)', norm(c['text']))
            if cm:
                d = abs(int(cm.group(1)) - w)
                if d < bd:
                    bd, best = d, c
        return best

    return None

def process_node(node, selected, facts, cfg):
    nid, ntext = node['id'], node['text']
    children = node.get('children', [])

    if children:
        sel = selected.get(nid)
        if sel is None:
            return []
        if ntext.rstrip().endswith(':'):
            return [(sel['id'], sel['text'])]
        if nid in cfg.get('combine_parents', []):
            return [(nid, ntext + ' ' + sel['text'])]
        return [(nid, ntext), (sel['id'], sel['text'])]

    if nid in cfg.get('fixed_ids', []):
        return [(nid, ntext)]
    if node.get('original_level') == 3:
        if score_orphan(ntext, facts) >= cfg.get('threshold', 0.25):
            return [(nid, ntext)]
        return []
    if simple_score(ntext, facts['raw_norm']) >= cfg.get('threshold', 0.25):
        return [(nid, ntext)]
    return []

def build_result(items, by_id, facts, cfg):
    selected = {}
    for it in items:
        if it.get('children'):
            selected[it['id']] = select_child(it, facts)

    result = []
    sections = sorted([it for it in items if it.get('level') == 1],
                      key=lambda x: natural_key(x['id']))
    choose_one = cfg.get('choose_one_sections', {})

    for section in sections:
        lines = []
        children = section.get('children', [])
        if section['id'] in choose_one:
            handler = choose_one[section['id']]
            if handler == 'complaint':
                sel = select_complaint(children, facts)
                if sel:
                    lines.append((sel['id'], sel['text']))
            else:
                best, bs = None, 0
                for c in children:
                    s = simple_score(c['text'], facts['raw_norm'])
                    if s > bs:
                        bs, best = s, c
                if best is None and children:
                    best = children[0]
                if best:
                    lines.append((best['id'], best['text']))
        else:
            for child in children:
                lines.extend(process_node(child, selected, facts, cfg))
        if lines:
            result.append((section['id'], section['text']))
            result.extend(lines)
    return result

# ============================================================
# ЧЕКЕР
# ============================================================
def check_classifier(path, cfg):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    seen, labels = set(), {}
    problems, warnings = [], []
    by_level = {1: 0, 2: 0, 3: 0}

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or len(row) < 2:
            continue
        code, text = row[0], row[1]
        if code is None:
            continue
        cs = str(code).strip()
        if cs.endswith('.0') and cs[:-2].isdigit():
            cs = cs[:-2]
        if not cs or cs.upper() == 'ID_CLASS':  # строка заголовка
            continue
        if cs in seen:
            problems.append(f"строка {i}: дубль {cs!r}")
            continue
        seen.add(cs)
        if re.fullmatch(r'\d', cs):
            by_level[1] += 1
        elif re.fullmatch(r'\d[a-zа-я]', cs):
            by_level[2] += 1
        elif re.fullmatch(r'\d[a-zа-я]\d+', cs):
            by_level[3] += 1
        else:
            problems.append(f"строка {i}: код {cs!r} невалиден")
            continue
        if not text or not str(text).strip():
            problems.append(f"строка {i}: {cs!r} без текста")
            continue
        labels.setdefault(str(text).strip(), []).append(cs)

    for lbl, codes in labels.items():
        if len(codes) > 1:
            warnings.append(f"дубли меток {codes}: {lbl[:60]!r}")

    for c in cfg.get('fixed_ids', []):
        if c not in seen:
            warnings.append(f"fixed_ids: {c!r} нет в xlsx")
    for c in cfg.get('combine_parents', []):
        if c not in seen:
            warnings.append(f"combine_parents: {c!r} нет в xlsx")

    return {'ok': not problems, 'total': len(seen),
            'by_level': by_level, 'problems': problems, 'warnings': warnings}

def format_check(r):
    out = [f"[{'OK' if r['ok'] else 'ПРОБЛЕМЫ'}] кодов: {r['total']}"]
    out.append(f"  секции {r['by_level'][1]}, поля {r['by_level'][2]}, опции {r['by_level'][3]}")
    if r['problems']:
        out.append(f"\nОшибки ({len(r['problems'])}):")
        out += [f"  ! {p}" for p in r['problems']]
    if r['warnings']:
        out.append(f"\nПредупреждения ({len(r['warnings'])}):")
        out += [f"  ? {w}" for w in r['warnings']]
    return "\n".join(out)

# ============================================================
# CLI
# ============================================================
def load_test_cases(path):
    """Эталоны из xlsx: C = ID_INPUT_RAW, D = INPUT_RAW, E = ID_CLASS_RESULT.
    Кейс начинается со строки, где заполнена C, и идёт до следующей такой строки."""
    wb = load_workbook(path, read_only=True, data_only=True)
    cases = []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        row = tuple(row) + (None,) * 6
        if row[2] is not None and str(row[2]).strip().upper() != 'ID_INPUT_RAW':
            cases.append({'id': row[2], 'raw': str(row[3] or ''), 'expected': []})
        elif cases and row[4] is not None:
            cs = str(row[4]).strip()
            if cs.endswith('.0'):
                cs = cs[:-2]
            cases[-1]['expected'].append(cs)
    return cases

def run_tests(path, cfg):
    items = load_classifier(path)
    by_id = build_hierarchy(items)
    failed = 0
    for case in load_test_cases(path):
        got = [i for i, _ in build_result(items, by_id, extract_facts(case['raw']), cfg)]
        exp = case['expected']
        missing = [i for i in exp if i not in got]
        extra = [i for i in got if i not in exp]
        ok = not missing and not extra
        failed += not ok
        print(f"[{'OK' if ok else 'FAIL'}] кейс {case['id']}: {case['raw'][:70]}")
        if missing:
            print(f"    нет в результате: {missing}")
        if extra:
            print(f"    лишнее:           {extra}")
    return 1 if failed else 0

def cli_generate(args, cfg):
    if not Path(args.xlsx).exists():
        print(f"Файл не найден: {args.xlsx}", file=sys.stderr)
        return 2
    raw = args.raw
    if args.raw_file:
        raw = Path(args.raw_file).read_text('utf-8')
    if not raw or not raw.strip():
        print("Пустой raw", file=sys.stderr)
        return 2

    items = load_classifier(args.xlsx)
    by_id = build_hierarchy(items)
    facts = extract_facts(raw)
    result = build_result(items, by_id, facts, cfg)
    for w in facts.get('warnings', []):
        print(f"⚠ {w}", file=sys.stderr)

    if args.format == 'doc':
        out = "\n".join(t for _, t in result)
    elif args.format == 'pipe':
        out = "\n".join(f"{i} | {t}" for i, t in result)
    else:
        out = "\n".join(f"{i}\t{t}" for i, t in result)

    if args.out:
        Path(args.out).write_text(out, encoding='utf-8-sig')
        print(f"Сохранено: {args.out}")
    else:
        print(out)
    return 0

# ============================================================
# GUI
# ============================================================
def gui_main(args, cfg):
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.title("Ортодонтический классификатор")
    root.geometry("1200x820")

    top = tk.Frame(root); top.pack(fill='x', padx=10, pady=6)
    tk.Label(top, text="order.xlsx:").pack(side='left')
    xlsx_var = tk.StringVar(value=args.xlsx or "order.xlsx")
    tk.Entry(top, textvariable=xlsx_var, width=60).pack(side='left', padx=4)
    tk.Button(top, text="Обзор", command=lambda: xlsx_var.set(
        filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")]) or xlsx_var.get()
    )).pack(side='left')

    state = {'items': None, 'by_id': None}
    info = tk.Label(top, text="не загружен", fg='gray')
    info.pack(side='left', padx=8)

    def do_load():
        try:
            items = load_classifier(xlsx_var.get())
            state['items'] = items
            state['by_id'] = build_hierarchy(items)
            info.config(text=f"загружено: {len(items)}", fg='green')
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))

    def do_check():
        try:
            r = check_classifier(xlsx_var.get(), cfg)
            win = tk.Toplevel(root); win.title("Проверка"); win.geometry("820x600")
            txt = tk.Text(win, wrap='word', font=('Consolas', 10))
            txt.pack(fill='both', expand=True, padx=10, pady=10)
            txt.insert('1.0', format_check(r))
            txt.config(state='disabled')
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))

    tk.Button(top, text="Загрузить", command=do_load).pack(side='left', padx=4)
    tk.Button(top, text="Проверить", command=do_check).pack(side='left', padx=4)

    tk.Label(root, text="Raw-запись врача:").pack(anchor='w', padx=10)
    raw_txt = tk.Text(root, height=8, wrap='word')
    raw_txt.pack(fill='x', padx=10, pady=4)
    raw_txt.insert('1.0',
        "активация, лигатурная, этап рабочий, класс по молярам 1 и клыкам, "
        "замена дуги на верху 18 25 нити, ношение эластиков перепел, явка через 4 недели.")

    btns = tk.Frame(root); btns.pack(fill='x', padx=10, pady=4)

    def do_generate():
        if not state['items']:
            do_load()
            if not state['items']:
                return
        raw = raw_txt.get('1.0', 'end').strip()
        facts = extract_facts(raw)
        result = build_result(state['items'], state['by_id'], facts, cfg)
        out_txt.delete('1.0', 'end')
        text = "\n".join(f"{i}\t{t}" for i, t in result)
        for w in facts.get('warnings', []):
            text += f"\n⚠ {w}"
        out_txt.insert('1.0', text)

    def do_save():
        text = out_txt.get('1.0', 'end').strip()
        if not text:
            return
        p = filedialog.asksaveasfilename(defaultextension=".txt",
                                         filetypes=[("Text", "*.txt")])
        if p:
            Path(p).write_text(text, encoding='utf-8-sig')

    tk.Button(btns, text="Сформировать", command=do_generate, height=2, width=18,
              bg='#4CAF50', fg='white').pack(side='left', padx=4)
    tk.Button(btns, text="Сохранить TXT", command=do_save, height=2, width=18)\
        .pack(side='left', padx=4)

    tk.Label(root, text="Результат  (ID_CLASS ⇥ TEXT):").pack(anchor='w', padx=10)
    out_txt = tk.Text(root, height=24, wrap='word')
    out_txt.pack(fill='both', expand=True, padx=10, pady=4)

    if Path(xlsx_var.get()).exists():
        do_load()
    root.mainloop()
    return 0

# ============================================================
# MAIN
# ============================================================
def main():
    ap = argparse.ArgumentParser(description="Ортодонтический маппинг raw → ID_CLASS")
    ap.add_argument('--xlsx', help='Путь к order.xlsx')
    ap.add_argument('--raw', help='Raw-текст')
    ap.add_argument('--raw-file', help='Файл с raw-текстом')
    ap.add_argument('--out', help='Куда сохранить результат')
    ap.add_argument('--format', choices=['tsv', 'pipe', 'doc'], default='tsv')
    ap.add_argument('--check', metavar='XLSX', help='Проверить классификатор')
    ap.add_argument('--test', metavar='XLSX', help='Сверить с эталонами (колонки C-F)')
    ap.add_argument('--config', help='Путь к ortodont_config.json')
    args = ap.parse_args()

    cfg = load_config(args.config)

    if args.test:
        return run_tests(args.test, cfg)

    if args.check:
        r = check_classifier(args.check, cfg)
        print(format_check(r))
        return 0 if r['ok'] else 1

    if args.raw or args.raw_file:
        return cli_generate(args, cfg)

    return gui_main(args, cfg)

if __name__ == '__main__':
    sys.exit(main() or 0)