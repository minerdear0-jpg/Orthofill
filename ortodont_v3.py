# -*- coding: utf-8 -*-
"""
ortodont.py — маппинг raw-записи врача на ID_CLASS классификатора.

Режимы:
  python ortodont.py                              # GUI
  python ortodont.py --check order.xlsx           # только проверка (exit 0/1)
  python ortodont.py --xlsx order.xlsx --raw "..." [--out result.txt]
  python ortodont.py --xlsx order.xlsx --raw-file raw.txt [--format tsv|pipe|doc]

Формат order.xlsx (лист 1):
  A = ID_CLASS  (1, 1a, 2b1, 4b10, ...)
  B = CLASS_LIST (текст классификатора)
"""
import sys, re, json, argparse
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Установите: pip install openpyxl")

# ============================================================
# КОНФИГ
# ============================================================
DEFAULT_CONFIG = {
    # Буалярные ID: выводятся всегда при активной секции
    "fixed_ids": ["2a", "2d", "3a", "3b", "3c", "3d",
                  "4a", "4j", "5a", "5b", "5c"],
    # ID, где текст родителя склеивается с текстом выбранного ребёнка
    "combine_parents": ["2b", "2c"],
    # Секции, где из детей выбирается один (по смыслу — жалобы)
    "choose_one_sections": {"1": "complaint"},
    # Порог нечёткого совпадения для «сиротских» и новых листовых позиций
    "threshold": 0.25,
}

def load_config(path=None):
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    p = Path(path) if path else Path(__file__).with_name("ortodont_config.json")
    if p.exists():
        try:
            data = json.loads(p.read_text("utf-8"))
            for k, v in data.items():
                if isinstance(v, dict) and isinstance(cfg.get(k), dict):
                    cfg[k].update(v)
                else:
                    cfg[k] = v
        except Exception as e:
            print(f"config: {e}", file=sys.stderr)
    return cfg

# ============================================================
# ТЕКСТ
# ============================================================
def norm(t):
    if t is None:
        return ''
    s = str(t).lower().replace('ё', 'е').replace('×', 'x')
    s = s.replace('–', '-').replace('—', '-')
    s = re.sub(r'<br\s*/?>', ' ', s, flags=re.I)
    return re.sub(r'\s+', ' ', s).strip()

def strip_html(t):
    if not t:
        return ''
    return re.sub(r'<br\s*/?>', ' ', str(t), flags=re.I).strip()

def natural_key(s):
    return [(0, int(x)) if x.isdigit() else (1, x)
            for x in re.split(r'(\d+)', s) if x]

# ============================================================
# ПАРСЕРЫ RAW
# ============================================================
def parse_system_type(n):
    if 'самолигир' in n:
        return 'self_ligating'
    if 'лигатур' in n:
        return 'ligature'
    return None

def parse_stage(n):
    if 'детализ' in n:
        return 'detail'
    if 'начальн' in n or 'нивелир' in n:
        return 'initial'
    return 'working'

def parse_classes(n):
    """Извлекает класс моляров и клыков. Клыки — только при явном числе рядом."""
    result = {'molar': None, 'canine': None}

    molar_pos = [m.start() for m in re.finditer(r'моляр', n)]
    canine_pos = [m.start() for m in re.finditer(r'клык', n)]

    for m in re.finditer(r'\b([1-3]|I{1,3})\b', n.upper()):
        v = m.group(1)
        if v in ('I', 'II', 'III'):
            val = {'I': 1, 'II': 2, 'III': 3}[v]
        else:
            val = int(v)
        pos = m.start()

        def nearest(lst):
            return min((abs(p - pos) for p in lst), default=999)

        md, cd = nearest(molar_pos), nearest(canine_pos)
        if md < cd and md < 30 and result['molar'] is None:
            result['molar'] = val
        elif cd < md and cd < 30 and result['canine'] is None:
            result['canine'] = val
    return result

def parse_wire_from_raw(n):
    if not re.search(r'(замен|смен)\w*\s+дуг|дуг\w*\s+(замен|смен)', n):
        return {}
    if 'обе' in n:
        arches = ['upper', 'lower']
    else:
        arches = []
        if 'верх' in n:
            arches.append('upper')
        if 'ниж' in n:
            arches.append('lower')
        if not arches:
            arches = ['upper', 'lower']
    size = None
    m = re.search(r'0?[.,]?0?(\d{2})\s*[x]\s*0?[.,]?0?(\d{2})', n)
    if m:
        size = (m.group(1), m.group(2))
    else:
        m = re.search(r'0?[.,]?0?(\d{2})(?:"|\s|$)', n)
        if m:
            size = (m.group(1), None)
    material = None
    if re.search(r'cuniti|cu[ -]?ni[ -]?ti', n):
        material = 'CuNiTi'
    elif re.search(r'\btma\b', n):
        material = 'TMA'
    elif re.search(r'\bss\b', n):
        material = 'SS'
    elif re.search(r'niti|ni[ -]?ti|нити|нитинол', n):
        material = 'NiTi'
    return {a: (size, material) for a in arches}

def parse_wire_from_label(label):
    n = norm(label)
    arch = 'upper' if 'верхн' in n else 'lower' if 'нижн' in n else None
    size = None
    m = re.search(r'0?[.,]?0?(\d{2})\s*x\s*0?[.,]?0?(\d{2})', n)
    if m:
        size = (m.group(1), m.group(2))
    else:
        m = re.search(r'0?[.,]?0?(\d{2})"', n)
        if m:
            size = (m.group(1), None)
    material = None
    if 'cuniti' in n or 'cu ni ti' in n or 'cu-ni-ti' in n:
        material = 'CuNiTi'
    elif re.search(r'\btma\b', n):
        material = 'TMA'
    elif re.search(r'\bss\b', n):
        material = 'SS'
    elif 'niti' in n or 'нити' in n:
        material = 'NiTi'
    return arch, size, material

def parse_elastics(n):
    for animal, key in [
        ('перепел', 'перепел'), ('перепёл', 'перепел'),
        ('лиса', 'лиса'), ('лису', 'лиса'),
        ('медвед', 'медведь'),
        ('антилоп', 'антилопа'),
        ('кролик', 'кролик'),
        ('зебра', 'зебра'), ('зебру', 'зебра'),
    ]:
        if animal in n:
            return key
    return None

def parse_weeks(n):
    m = re.search(r'через\s+(\d+)\s*(?:недел|нед)', n)
    if m:
        return int(m.group(1))
    m = re.search(r'(\d+)\s*(?:недел|нед)', n)
    if m:
        return int(m.group(1))
    return 4

def parse_hygiene(n):
    return 'unsatisfactory' if 'гигиен' in n and 'неудовлетвор' in n else 'satisfactory'

def parse_mucosa(n):
    return 'erosion' if 'эрози' in n else 'normal'

def parse_bracket_state(n):
    return 'unsatisfactory' if 'брекет' in n and 'неудовлетвор' in n else 'satisfactory'

def parse_dynamics(n):
    return 'no_dynamics' if 'без' in n and ('динамик' in n or 'существен' in n) else 'plan'

def parse_ligature_change(n, st):
    if re.search(r'(смен|замен)\w*\s+.*лигатур|лигатур\w*\s+(смен|замен)', n):
        return True
    return st == 'ligature'

def parse_change_arches(n):
    if 'обе' in n:
        return 'both'
    u, l = 'верх' in n, 'ниж' in n
    if u and l:
        return 'both'
    return 'upper' if u else 'lower' if l else 'both'

def parse_wire_adjustments(n):
    out = []
    if re.search(r'изгиб\w*\s+перв|перв\w*\s+порядк', n):
        out.append(1)
    if re.search(r'изгиб\w*\s+втор|втор\w*\s+порядк', n):
        out.append(2)
    if re.search(r'изгиб\w*\s+трет|трет\w*\s+порядк', n):
        out.append(3)
    return out

def extract_facts(raw):
    n = norm(raw)
    st = parse_system_type(n)
    return {
        'raw_norm': n,
        'system_type': st,
        'stage': parse_stage(n),
        'classes': parse_classes(n),
        'wires': parse_wire_from_raw(n),
        'elastics': parse_elastics(n),
        'weeks': parse_weeks(n),
        'hygiene': parse_hygiene(n),
        'mucosa': parse_mucosa(n),
        'bracket_state': parse_bracket_state(n),
        'dynamics': parse_dynamics(n),
        'ligature_change': parse_ligature_change(n, st),
        'change_arches': parse_change_arches(n),
        'wire_adjustments': parse_wire_adjustments(n),
        'has_binding': 'подвяз' in n,
    }

# ============================================================
# КЛАССИФИКАТОР
# ============================================================
def load_classifier(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    items = []
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        code, text = row[0], row[1]
        if code is None or text is None:
            continue
        cs = str(code).strip()
        if cs.endswith('.0') and cs[:-2].isdigit():
            cs = cs[:-2]
        if not re.fullmatch(r'\d|\d[a-zа-я]|\d[a-zа-я]\d+', cs):
            continue
        txt = strip_html(text)
        if not txt:
            continue
        items.append({'id': cs, 'text': txt})
    return items

def build_hierarchy(items):
    by_id = {it['id']: it for it in items}
    for it in items:
        cs = it['id']
        if re.fullmatch(r'\d', cs):
            it['level'], it['parent'] = 1, None
        elif re.fullmatch(r'\d[a-zа-я]', cs):
            it['level'], it['parent'] = 2, cs[0]
        elif re.fullmatch(r'\d[a-zа-я]\d+', cs):
            it['level'] = 3
            it['parent'] = cs[:re.search(r'\d+$', cs).start()]
        else:
            it['level'], it['parent'] = 0, None
        it['original_level'] = it['level']
        it['children'] = []

    # Осиротевшие level-3 → повышаем до level-2 (самодостаточные)
    for it in items:
        if it['level'] == 3 and it['parent'] not in by_id:
            it['level'] = 2
            it['parent'] = it['id'][0]

    for it in items:
        if it['parent'] and it['parent'] in by_id:
            by_id[it['parent']]['children'].append(it)
    for it in items:
        it['children'].sort(key=lambda x: natural_key(x['id']))
    return by_id

# ============================================================
# МАТЧИНГ
# ============================================================
def match_class_num(text, num):
    m = re.search(r'\b(I{1,4}|[1-4])\s*класс', norm(text).upper())
    if not m:
        return False
    v = m.group(1)
    if v in ('I', 'II', 'III', 'IV'):
        return {'I': 1, 'II': 2, 'III': 3, 'IV': 4}[v] == num
    return int(v) == num

def match_order(text, order):
    return {1: 'перв', 2: 'втор', 3: 'трет'}[order] in norm(text)

def simple_score(text, raw_norm):
    kws = [w for w in re.findall(r'[а-яa-z]{3,}|\d+', norm(text))]
    if not kws:
        return 0
    return sum(1 for k in kws if k in raw_norm) / len(kws)

def score_orphan(text, facts):
    n = norm(text)
    st = facts['system_type'] or 'ligature'
    if 'фиксаци' in n and 'лигатур' in n:
        return 1.0 if st != 'self_ligating' else 0.0
    if 'клипс' in n:
        return 1.0 if st == 'self_ligating' else 0.0
    return simple_score(text, facts['raw_norm'])

def select_complaint(children, facts):
    n = facts['raw_norm']
    for c in children:
        ct = norm(c['text'])
        if 'отклеива' in ct and re.search(r'отклеива|отвалил', n):
            return c
        if 'выход' in ct and 'дуг' in ct and re.search(r'выход\s+дуги', n):
            return c
        if 'перелом' in ct and re.search(r'перелом\s+дуг|сломал\w*\s+дуг', n):
            return c
    return children[0] if children else None

def select_child(parent, facts):
    pid = parent['id']
    children = parent['children']
    if not children:
        return None

    if pid == '2b':
        st = facts['system_type']
        if st is None:
            return None
        for c in children:
            ct = norm(c['text'])
            if st == 'self_ligating' and 'самолигир' in ct:
                return c
            if st == 'ligature' and 'лигатур' in ct and 'самолигир' not in ct:
                return c
        return None

    if pid == '2c':
        stg = facts['stage']
        for c in children:
            ct = norm(c['text'])
            if stg == 'detail' and 'детализ' in ct:
                return c
            if stg == 'initial' and ('начальн' in ct or 'нивелир' in ct):
                return c
            if stg == 'working' and 'рабоч' in ct:
                return c
        return None

    if pid == '3e':
        if facts['mucosa'] == 'erosion':
            for c in children:
                if 'эрози' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'без видимых' in ct or 'норм' in ct:
                return c
        return children[0]

    if pid == '3f':
        if facts['hygiene'] == 'unsatisfactory':
            for c in children:
                if 'неудовлетвор' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'удовлетвор' in ct and 'неудовлетвор' not in ct:
                return c
        return children[0]

    if pid in ('3g', '3h', '3y', '3j'):
        tooth = 'моляр' if pid in ('3g', '3h') else 'клык'
        cls = facts['classes']['molar' if tooth == 'моляр' else 'canine']
        if cls is None:
            return None
        side = 'справа' if pid in ('3g', '3y') else 'слева'
        for c in children:
            ct = norm(c['text'])
            if side in ct and tooth[:5] in ct and match_class_num(ct, cls):
                return c
        return None

    if pid == '3k':
        if facts['bracket_state'] == 'unsatisfactory':
            for c in children:
                if 'неудовлетвор' in norm(c['text']):
                    return c
        for c in children:
            ct = norm(c['text'])
            if 'удовлетвор' in ct and 'неудовлетвор' not in ct:
                return c
        return children[0]

    if pid in ('4b', '4c'):
        arch = 'upper' if pid == '4b' else 'lower'
        wire_info = facts['wires'].get(arch)
        if wire_info is None:
            return None
        size, material = wire_info
        if size is None:
            return None
        for c in children:
            la, ls, lm = parse_wire_from_label(c['text'])
            if la != arch or ls != size:
                continue
            if material and lm and material != lm:
                continue
            return c
        return None

    if pid == '4d':
        if not facts['ligature_change']:
            return None
        arches = facts['change_arches']
        for c in children:
            ct = norm(c['text'])
            if arches == 'both' and 'обе' in ct:
                return c
            if arches == 'upper' and 'верхн' in ct:
                return c
            if arches == 'lower' and 'нижн' in ct:
                return c
        return None

    if pid == '4e':
        if not facts['has_binding']:
            return None
        arches = facts['change_arches']
        for c in children:
            ct = norm(c['text'])
            if arches == 'upper' and 'верхн' in ct:
                return c
            if arches == 'lower' and 'нижн' in ct:
                return c
        return None

    if pid == '4f':
        el = facts['elastics']
        if el is None:
            return None
        for c in children:
            if el in norm(c['text']):
                return c
        return None

    if pid == '4g':
        if not facts['wire_adjustments']:
            return None
        for order in facts['wire_adjustments']:
            for c in children:
                if match_order(c['text'], order):
                    return c
        return None

    if pid == '5d':
        w = facts['weeks']
        for c in children:
            if str(w) in norm(c['text']):
                return c
        best, bd = None, 999
        for c in children:
            cm = re.search(r'(\d+)', norm(c['text']))
            if cm:
                d = abs(int(cm.group(1)) - w)
                if d < bd:
                    bd, best = d, c
        return best

    return None

def process_node(node, selected, facts, cfg):
    nid, ntext = node['id'], node['text']
    children = node.get('children', [])

    if children:
        sel = selected.get(nid)
        if sel is None:
            return []
        if ntext.rstrip().endswith(':'):
            return [(sel['id'], sel['text'])]
        if nid in cfg.get('combine_parents', []):
            return [(nid, ntext + ' ' + sel['text'])]
        return [(nid, ntext), (sel['id'], sel['text'])]

    if nid in cfg.get('fixed_ids', []):
        return [(nid, ntext)]
    if node.get('original_level') == 3:
        if score_orphan(ntext, facts) >= cfg.get('threshold', 0.25):
            return [(nid, ntext)]
        return []
    if simple_score(ntext, facts['raw_norm']) >= cfg.get('threshold', 0.25):
        return [(nid, ntext)]
    return []

def build_result(items, by_id, facts, cfg):
    selected = {}
    for it in items:
        if it.get('children'):
            selected[it['id']] = select_child(it, facts)

    result = []
    sections = sorted([it for it in items if it.get('level') == 1],
                      key=lambda x: natural_key(x['id']))
    choose_one = cfg.get('choose_one_sections', {})

    for section in sections:
        lines = []
        children = section.get('children', [])
        if section['id'] in choose_one:
            handler = choose_one[section['id']]
            if handler == 'complaint':
                sel = select_complaint(children, facts)
                if sel:
                    lines.append((sel['id'], sel['text']))
            else:
                best, bs = None, 0
                for c in children:
                    s = simple_score(c['text'], facts['raw_norm'])
                    if s > bs:
                        bs, best = s, c
                if best is None and children:
                    best = children[0]
                if best:
                    lines.append((best['id'], best['text']))
        else:
            for child in children:
                lines.extend(process_node(child, selected, facts, cfg))
        if lines:
            result.append((section['id'], section['text']))
            result.extend(lines)
    return result

# ============================================================
# ЧЕКЕР
# ============================================================
def check_classifier(path, cfg):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    seen, labels = set(), {}
    problems, warnings = [], []
    by_level = {1: 0, 2: 0, 3: 0}

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or len(row) < 2:
            continue
        code, text = row[0], row[1]
        if code is None:
            continue
        cs = str(code).strip()
        if cs.endswith('.0') and cs[:-2].isdigit():
            cs = cs[:-2]
        if not cs:
            continue
        if cs in seen:
            problems.append(f"строка {i}: дубль {cs!r}")
            continue
        seen.add(cs)
        if re.fullmatch(r'\d', cs):
            by_level[1] += 1
        elif re.fullmatch(r'\d[a-zа-я]', cs):
            by_level[2] += 1
        elif re.fullmatch(r'\d[a-zа-я]\d+', cs):
            by_level[3] += 1
        else:
            problems.append(f"строка {i}: код {cs!r} невалиден")
            continue
        if not text or not str(text).strip():
            problems.append(f"строка {i}: {cs!r} без текста")
            continue
        labels.setdefault(str(text).strip(), []).append(cs)

    for lbl, codes in labels.items():
        if len(codes) > 1:
            warnings.append(f"дубли меток {codes}: {lbl[:60]!r}")

    for c in cfg.get('fixed_ids', []):
        if c not in seen:
            warnings.append(f"fixed_ids: {c!r} нет в xlsx")
    for c in cfg.get('combine_parents', []):
        if c not in seen:
            warnings.append(f"combine_parents: {c!r} нет в xlsx")

    return {'ok': not problems, 'total': len(seen),
            'by_level': by_level, 'problems': problems, 'warnings': warnings}

def format_check(r):
    out = [f"[{'OK' if r['ok'] else 'ПРОБЛЕМЫ'}] кодов: {r['total']}"]
    out.append(f"  секции {r['by_level'][1]}, поля {r['by_level'][2]}, опции {r['by_level'][3]}")
    if r['problems']:
        out.append(f"\nОшибки ({len(r['problems'])}):")
        out += [f"  ! {p}" for p in r['problems']]
    if r['warnings']:
        out.append(f"\nПредупреждения ({len(r['warnings'])}):")
        out += [f"  ? {w}" for w in r['warnings']]
    return "\n".join(out)

# ============================================================
# CLI
# ============================================================
def cli_generate(args, cfg):
    if not Path(args.xlsx).exists():
        print(f"Файл не найден: {args.xlsx}", file=sys.stderr)
        return 2
    raw = args.raw
    if args.raw_file:
        raw = Path(args.raw_file).read_text('utf-8')
    if not raw or not raw.strip():
        print("Пустой raw", file=sys.stderr)
        return 2

    items = load_classifier(args.xlsx)
    by_id = build_hierarchy(items)
    facts = extract_facts(raw)
    result = build_result(items, by_id, facts, cfg)

    if args.format == 'doc':
        out = "\n".join(t for _, t in result)
    elif args.format == 'pipe':
        out = "\n".join(f"{i} | {t}" for i, t in result)
    else:
        out = "\n".join(f"{i}\t{t}" for i, t in result)

    if args.out:
        Path(args.out).write_text(out, encoding='utf-8-sig')
        print(f"Сохранено: {args.out}")
    else:
        print(out)
    return 0

# ============================================================
# GUI
# ============================================================
def gui_main(args, cfg):
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.title("Ортодонтический классификатор")
    root.geometry("1200x820")

    top = tk.Frame(root); top.pack(fill='x', padx=10, pady=6)
    tk.Label(top, text="order.xlsx:").pack(side='left')
    xlsx_var = tk.StringVar(value=args.xlsx or "order.xlsx")
    tk.Entry(top, textvariable=xlsx_var, width=60).pack(side='left', padx=4)
    tk.Button(top, text="Обзор", command=lambda: xlsx_var.set(
        filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")]) or xlsx_var.get()
    )).pack(side='left')

    state = {'items': None, 'by_id': None}
    info = tk.Label(top, text="не загружен", fg='gray')
    info.pack(side='left', padx=8)

    def do_load():
        try:
            items = load_classifier(xlsx_var.get())
            state['items'] = items
            state['by_id'] = build_hierarchy(items)
            info.config(text=f"загружено: {len(items)}", fg='green')
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))

    def do_check():
        try:
            r = check_classifier(xlsx_var.get(), cfg)
            win = tk.Toplevel(root); win.title("Проверка"); win.geometry("820x600")
            txt = tk.Text(win, wrap='word', font=('Consolas', 10))
            txt.pack(fill='both', expand=True, padx=10, pady=10)
            txt.insert('1.0', format_check(r))
            txt.config(state='disabled')
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))

    tk.Button(top, text="Загрузить", command=do_load).pack(side='left', padx=4)
    tk.Button(top, text="Проверить", command=do_check).pack(side='left', padx=4)

    tk.Label(root, text="Raw-запись врача:").pack(anchor='w', padx=10)
    raw_txt = tk.Text(root, height=8, wrap='word')
    raw_txt.pack(fill='x', padx=10, pady=4)
    raw_txt.insert('1.0',
        "активация, лигатурная, этап рабочий, класс по молярам 1 и клыкам, "
        "замена дуги на верху 18 25 нити, ношение эластиков перепел, явка через 4 недели.")

    btns = tk.Frame(root); btns.pack(fill='x', padx=10, pady=4)

    def do_generate():
        if not state['items']:
            do_load()
            if not state['items']:
                return
        raw = raw_txt.get('1.0', 'end').strip()
        facts = extract_facts(raw)
        result = build_result(state['items'], state['by_id'], facts, cfg)
        out_txt.delete('1.0', 'end')
        out_txt.insert('1.0', "\n".join(f"{i}\t{t}" for i, t in result))

    def do_save():
        text = out_txt.get('1.0', 'end').strip()
        if not text:
            return
        p = filedialog.asksaveasfilename(defaultextension=".txt",
                                         filetypes=[("Text", "*.txt")])
        if p:
            Path(p).write_text(text, encoding='utf-8-sig')

    tk.Button(btns, text="Сформировать", command=do_generate, height=2, width=18,
              bg='#4CAF50', fg='white').pack(side='left', padx=4)
    tk.Button(btns, text="Сохранить TXT", command=do_save, height=2, width=18)\
        .pack(side='left', padx=4)

    tk.Label(root, text="Результат  (ID_CLASS ⇥ TEXT):").pack(anchor='w', padx=10)
    out_txt = tk.Text(root, height=24, wrap='word')
    out_txt.pack(fill='both', expand=True, padx=10, pady=4)

    if Path(xlsx_var.get()).exists():
        do_load()
    root.mainloop()
    return 0

# ============================================================
# MAIN
# ============================================================
def main():
    ap = argparse.ArgumentParser(description="Ортодонтический маппинг raw → ID_CLASS")
    ap.add_argument('--xlsx', help='Путь к order.xlsx')
    ap.add_argument('--raw', help='Raw-текст')
    ap.add_argument('--raw-file', help='Файл с raw-текстом')
    ap.add_argument('--out', help='Куда сохранить результат')
    ap.add_argument('--format', choices=['tsv', 'pipe', 'doc'], default='tsv')
    ap.add_argument('--check', metavar='XLSX', help='Проверить классификатор')
    ap.add_argument('--config', help='Путь к ortodont_config.json')
    args = ap.parse_args()

    cfg = load_config(args.config)

    if args.check:
        r = check_classifier(args.check, cfg)
        print(format_check(r))
        return 0 if r['ok'] else 1

    if args.raw or args.raw_file:
        return cli_generate(args, cfg)

    return gui_main(args, cfg)

if __name__ == '__main__':
    sys.exit(main() or 0)