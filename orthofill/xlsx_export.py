"""Пресет → xlsx того же формата, что читает load_workbook_file:
первый лист A–F (классификатор и эталоны), «Настройки», «Словарь»."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as XlBook

from .methods import METHOD_LABELS, WHEN_LABELS
from .logic import format as format_condition
from .workbook import (COMBINE_LABELS, DICTIONARY_HEADER, DICTIONARY_SHEET, HEADER, PHRASES_HEADER,
                       PHRASES_SHEET, SETTINGS_HEADER,
                       SETTINGS_SHEET, Workbook)


def _num(code: str):
    """«1» → 1 (как в исходных файлах врача), «4a» остаётся строкой."""
    return int(code) if code.isdigit() else code


def export_workbook(wb: Workbook, path: Path) -> None:
    """PermissionError (файл открыт в Excel) пробрасывается вызывающему."""
    xl = XlBook()
    ws = xl.active
    ws.title = 'Пресет'
    ws.append([HEADER, 'CLASS_LIST', 'ID_INPUT_RAW', 'INPUT_RAW', 'ID_CLASS_RESULT', 'RESULT'])
    ref_rows: list[tuple] = []
    for case in wb.references:
        ref_rows.append((_num(case.case_id), case.raw, None, None))
        ref_rows += [(None, None, _num(code) if code else None, text) for code, text in case.lines]
    for i in range(max(len(wb.items), len(ref_rows))):
        a = (_num(wb.items[i].id), wb.items[i].text) if i < len(wb.items) else (None, None)
        ws.append(list(a) + list(ref_rows[i] if i < len(ref_rows) else (None,) * 4))

    st = xl.create_sheet(SETTINGS_SHEET)
    st.append(SETTINGS_HEADER)
    partners: dict[str, list[str]] = {}
    for a, b in wb.not_together:            # пара — в обеих строках
        partners.setdefault(a, []).append(b)
        partners.setdefault(b, []).append(a)
    for it in wb.items:
        method = METHOD_LABELS.get(it.method) if it.children and it.method else None
        when = WHEN_LABELS.get(it.when) if it.when else None
        combine = COMBINE_LABELS.get(it.combine) if it.combine else None
        cond = format_condition(it.condition) or None
        apart = ', '.join(partners.get(it.id, [])) or None
        if method or when or it.default or it.words or combine or cond or apart:
            st.append([_num(it.id), method, when, it.default, ', '.join(it.words) or None,
                       combine, cond, apart])

    ft = xl.create_sheet(PHRASES_SHEET)
    ft.append(PHRASES_HEADER)
    for ph in wb.phrases:
        ft.append([ph.name, ph.template, ph.uses])
    ft.column_dimensions['A'].width = 16
    ft.column_dimensions['B'].width = 70
    ft.column_dimensions['C'].width = 14

    dt = xl.create_sheet(DICTIONARY_SHEET)
    dt.append(DICTIONARY_HEADER)
    for e in wb.dictionary:
        dt.append([e.word, e.item_text, e.code or None, e.added or None, e.hits])

    for sheet, widths in ((ws, (10, 60, 12, 50, 14, 60)), (st, (8, 34, 18, 14, 30, 12, 34, 16)),
                          (dt, (18, 50, 8, 14, 14))):
        for col, w in zip('ABCDEFGH', widths):
            sheet.column_dimensions[col].width = w
    xl.save(path)
