"""Файл врача (xlsx): классификатор (A–B) и эталоны (C–F)."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .methods import METHOD_LABELS, WHEN_LABELS

HEADER = 'ID_CLASS'
SETTINGS_SHEET = 'Настройки'
DICTIONARY_SHEET = 'Словарь'
PHRASES_SHEET = 'Фразы'
PHRASES_HEADER = ['Имя', 'Шаблон', 'Использовано']
SETTINGS_HEADER = ['Код', 'Способ выбора', 'Когда в протоколе', 'По умолчанию', 'Слова',
                   'Сочетание', 'Условие', 'Не вместе с']
COMBINE_LABELS = {'and': 'И'}       # пусто или «ИЛИ» — один вариант
DICTIONARY_HEADER = ['Слово', 'Пункт', 'Код', 'Добавлено', 'Сработало раз']
_CODE_RE = re.compile(r'\d|\d[a-zа-я]|\d[a-zа-я]\d+')


class WorkbookError(Exception):
    """Файл нельзя использовать: нет, не xlsx, нет строки заголовка."""


@dataclass
class Item:
    id: str
    text: str
    level: int                      # 1 раздел, 2 пункт, 3 вариант
    parent: str | None
    orphan: bool = False            # вариант без своей группы (4h1)
    children: list['Item'] = field(default_factory=list)
    # Свойства пресета (None — угадать, см. engine.guess_settings)
    method: str | None = None       # группа: способ выбора варианта (ключ METHOD_LABELS)
    when: str | None = None         # отдельный пункт: 'always' | 'mentioned'
    default: str | None = None      # группа: код варианта, если врач не написал
    words: list[str] = field(default_factory=list)   # доп. слова упоминания
    combine: str | None = None      # группа: 'and' — все упомянутые варианты; None — один (ИЛИ)
    condition: dict | None = None   # пункт: условие (orthofill/logic.py), None — нет


@dataclass
class ReferenceCase:
    case_id: str
    raw: str
    lines: list[tuple[str, str]]    # (код или '', текст) как в E/F


@dataclass
class DictEntry:
    """Слово врача → пункт пресета (по тексту пункта; код — подсказка)."""
    word: str
    item_text: str
    code: str
    added: str = ''
    hits: int = 0


@dataclass
class Phrase:
    """Фраза врача: `.имя` → шаблон с полями (orthofill/phrases.py)."""
    name: str
    template: str
    uses: int = 0


@dataclass
class Workbook:
    path: Path | None
    sha256: str
    items: list[Item]
    by_id: dict[str, Item]
    sections: list[Item]
    references: list[ReferenceCase]
    problems: list[str]             # с ними работать нельзя
    warnings: list[str]
    name: str = ''                  # имя пресета
    source: str = ''                # файл, из которого пресет импортирован
    dictionary: list[DictEntry] = field(default_factory=list)
    not_together: list[tuple[str, str]] = field(default_factory=list)   # пары (a, b), a < b
    phrases: list[Phrase] = field(default_factory=list)


def clean_code(value) -> str:
    s = str(value).strip().lower()
    if re.fullmatch(r'\d+\.0', s):
        s = s[:-2]
    return s


def clean_text(value) -> str:
    if value is None:
        return ''
    return re.sub(r'<br\s*/?>', ' ', str(value), flags=re.I).strip()


def _open_book(path: Path):
    if not path.exists():
        raise WorkbookError(f'Файл не найден: {path}')
    try:
        return load_workbook(path, data_only=True)
    except Exception as e:  # повреждён / не xlsx
        raise WorkbookError(f'Не удалось открыть {path.name}: {e}') from e


def find_header_row(ws) -> int:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip().upper() == HEADER:
                if cell.column == 1:
                    return cell.row
                raise WorkbookError(
                    f'Строка заголовка {HEADER} найдена в колонке {cell.column_letter}, '
                    f'а должна быть в A (строка {cell.row})')
    raise WorkbookError(f'Не найдена строка заголовка: в колонке A должно стоять {HEADER} '
                        f'(проверьте первые строки файла)')


def _row_values(ws, header_row):
    for row in ws.iter_rows(min_row=header_row + 1, max_col=6, values_only=True):
        yield tuple(row) + (None,) * (6 - len(row))


def _level(code: str) -> int:
    return 1 if len(code) == 1 else 2 if len(code) == 2 else 3


def load_workbook_file(path: str | Path) -> Workbook:
    path = Path(path)
    book = _open_book(path)
    ws = book.worksheets[0]
    header = find_header_row(ws)

    items: list[Item] = []
    by_id: dict[str, Item] = {}
    problems: list[str] = []
    warnings: list[str] = []
    references: list[ReferenceCase] = []

    for rownum, (a, b, c, d, e, f) in enumerate(_row_values(ws, header), start=header + 1):
        # --- классификатор
        if a is not None and str(a).strip():
            code = clean_code(a)
            text = clean_text(b)
            if not _CODE_RE.fullmatch(code):
                problems.append(f'строка {rownum}: код «{code}» не по формату (1, 1a, 1a1)')
            elif code in by_id:
                problems.append(f'строка {rownum}: код «{code}» повторяется')
            elif not text:
                problems.append(f'строка {rownum}: у кода «{code}» нет текста')
            else:
                lvl = _level(code)
                parent = None if lvl == 1 else code[:2] if lvl == 3 else code[0]
                it = Item(code, text, lvl, parent)
                items.append(it)
                by_id[code] = it
        # --- эталоны
        if c is not None and str(c).strip():
            case_id = clean_code(c)
            if any(r.case_id == case_id for r in references):
                problems.append(f'строка {rownum}: номер эталона «{case_id}» повторяется '
                                f'(колонка C) — дайте эталону новый номер')
            references.append(ReferenceCase(case_id, clean_text(d), []))
        elif references and f is not None and str(f).strip():
            code = clean_code(e) if e is not None and str(e).strip() else ''
            references[-1].lines.append((code, clean_text(f)))

    # Вариант без своей группы — самостоятельный пункт раздела
    for it in items:
        if it.level == 3 and it.parent not in by_id:
            it.level, it.parent, it.orphan = 2, it.id[0], True
    for it in items:
        if it.parent in by_id:
            by_id[it.parent].children.append(it)
        elif it.parent is not None:
            problems.append(f'код «{it.id}»: нет раздела «{it.parent}»')

    if SETTINGS_SHEET in book.sheetnames:
        found, notes, pairs = _read_settings(book[SETTINGS_SHEET], by_id)
        problems += found
        warnings += notes
    else:
        pairs = []
    dictionary = (_read_dictionary(book[DICTIONARY_SHEET])
                  if DICTIONARY_SHEET in book.sheetnames else [])
    phrases: list[Phrase] = []
    if PHRASES_SHEET in book.sheetnames:
        phrases, found = _read_phrases(book[PHRASES_SHEET])
        problems += found

    wb = Workbook(
        path=path,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        items=items,
        by_id=by_id,
        sections=[it for it in items if it.level == 1],
        references=references,
        problems=problems,
        warnings=warnings,
        name=path.stem,
        source=path.name,
        dictionary=dictionary,
        not_together=pairs,
        phrases=phrases,
    )
    normalize_references(wb)
    return wb


def _split_words(value) -> list[str]:
    return [w.strip() for w in re.split(r'[,;\n]', str(value or '')) if w.strip()]


def _read_settings(ws, by_id: dict[str, Item]) -> tuple[list[str], list[str], list[tuple[str, str]]]:
    """Лист «Настройки»: Код | Способ выбора | Когда в протоколе | По умолчанию |
    Слова | Сочетание | Условие | Не вместе с → (ошибки, предупреждения, пары)."""
    from .logic import parse
    method_of = {v: k for k, v in METHOD_LABELS.items()}
    when_of = {v: k for k, v in WHEN_LABELS.items()}
    problems: list[str] = []
    listed: dict[str, set[str]] = {}
    for rownum, row in enumerate(ws.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        code, method, when, default, words, combine, cond, apart = tuple(row) + (None,) * (8 - len(row))
        if code is None or not str(code).strip():
            continue
        where = f'{SETTINGS_SHEET}, строка {rownum}'
        it = by_id.get(clean_code(code))
        if it is None:
            problems.append(f'{where}: пункта «{clean_code(code)}» нет в классификаторе')
            continue
        if method:
            if str(method).strip() not in method_of:
                problems.append(f'{where}: способ выбора «{method}» не знаю')
            else:
                it.method = method_of[str(method).strip()]
        if when:
            if str(when).strip() not in when_of:
                problems.append(f'{where}: «{when}» — нужно «всегда», «если упомянут» или «по условию»')
            else:
                it.when = when_of[str(when).strip()]
        if default:
            it.default = clean_code(default)
        it.words = _split_words(words)
        cell = lambda col: f'{SETTINGS_SHEET}!{col}{rownum}'   # noqa: E731
        if combine is not None and str(combine).strip():
            c = str(combine).strip().upper()
            if c not in ('И', 'ИЛИ'):
                problems.append(f'{cell("F")}: «{combine}» — нужно «И» или «ИЛИ»')
            else:
                it.combine = 'and' if c == 'И' else None
        if cond is not None and str(cond).strip():
            try:
                it.condition = parse(str(cond), set(by_id))
            except ValueError as e:
                problems.append(f'{cell("G")}: {e}')
        for other in _split_words(apart):
            o = clean_code(other)
            if o not in by_id or o == it.id:
                problems.append(f'{cell("H")}: пункта «{o}» нет в классификаторе' if o != it.id
                                else f'{cell("H")}: пункт не может быть «не вместе» сам с собой')
            else:
                listed.setdefault(it.id, set()).add(o)
    pairs, notes = set(), []
    for a, others in listed.items():
        for b in others:
            pairs.add(tuple(sorted((a, b))))
            if a not in listed.get(b, set()):
                notes.append(f'пара «не вместе» {a} ↔ {b} дополнена (была только у {a})')
    return problems, notes, sorted(pairs)


def _read_phrases(ws) -> tuple[list[Phrase], list[str]]:
    """Лист «Фразы»: Имя | Шаблон | Использовано."""
    from .phrases import parse_template
    out, problems = [], []
    for rownum, row in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        name, template, uses = tuple(row) + (None,) * (3 - len(row))
        if name is None or not str(name).strip():
            continue
        name = str(name).strip().lstrip('.')
        try:
            parse_template(str(template or ''))
        except ValueError as e:
            problems.append(f'{PHRASES_SHEET}!B{rownum}: {e}')
            continue
        if not str(template or '').strip():
            problems.append(f'{PHRASES_SHEET}!B{rownum}: у фразы «{name}» нет шаблона')
            continue
        out.append(Phrase(name, str(template).strip(), int(uses or 0)))
    return out, problems


def _read_dictionary(ws) -> list[DictEntry]:
    out = []
    for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        word, item_text, code, added, hits = tuple(row) + (None,) * (5 - len(row))
        if word is None or not str(word).strip():
            continue
        out.append(DictEntry(str(word).strip(), clean_text(item_text),
                             clean_code(code) if code is not None else '',
                             str(added or ''), int(hits or 0)))
    return out


# ---------- «Сохранить как эталон» ----------
def _ws_key(text: str) -> str:
    return ' '.join(text.split())


def _tail_words(text: str, group: Item) -> set[str]:
    """Слова текста без начальных слов группы: «Тип системы: вестибулярная
    лигатурная.» → {вестибулярная, лигатурная}."""
    from .text import norm
    words = re.findall(r'[а-яa-z0-9]+', norm(text))
    head = re.findall(r'[а-яa-z0-9]+', norm(group.text))
    if words[:len(head)] == head:
        words = words[len(head):]
    return set(words)


def _same_group(prev_code: str, group: Item, wb: Workbook) -> bool:
    """Предыдущая строка — эта же группа или её вариант (строку группы не повторять)."""
    prev = wb.by_id.get(prev_code)
    return prev_code == group.id or (prev is not None and prev.parent == group.id
                                     and prev.level == 3)


def _variant_join_text(group: Item, variant: Item) -> list[str]:
    from .render import join_group
    return [_ws_key(join_group(group.text, variant.text)),
            _ws_key(f'{group.text} {variant.text}')]


def normalize_references(wb: Workbook) -> None:
    """Эталоны 0.1 → вид «по пунктам» (идемпотентно):
    «2b | Тип системы вестибулярная лигатурная.» → группа 2b + вариант 2b2;
    вариант без строки своей группы выше → строка группы вставляется."""
    for case in wb.references:
        out: list[tuple[str, str]] = []
        for code, text in case.lines:
            it = wb.by_id.get(code) if code else None
            if it is not None and it.children and _ws_key(text) != _ws_key(it.text):
                hit = next((c for c in it.children
                            if _ws_key(text) in _variant_join_text(it, c)), None)
                if hit is None:      # варианты переписаны целыми фразами — по словам после группы
                    tail = _tail_words(text, it)
                    cands = [c for c in it.children if tail and tail <= _tail_words(c.text, it)]
                    hit = cands[0] if len(cands) == 1 else None
                if hit is not None:
                    if not (out and _same_group(out[-1][0], it, wb)):
                        out.append((it.id, it.text))
                    out.append((hit.id, hit.text))
                    continue
            elif it is not None and it.children:
                if out and _same_group(out[-1][0], it, wb):
                    continue         # «группа, вариант, группа, вариант» → одна строка группы
            elif it is not None and it.level == 3 and not it.orphan and it.parent in wb.by_id:
                group = wb.by_id[it.parent]
                if not (out and _same_group(out[-1][0], group, wb)):
                    out.append((group.id, group.text))
            out.append((code, text))
        case.lines = out


def lines_for_reference(wb: Workbook, generated, final_text: str) -> list[tuple[str, str]]:
    """Строки, которые врач оставил в окне, → (код, текст пункта) для эталона.

    Строка совпала с выданной строкой → коды этой строки (группа и вариант
    отдельными строками); с текстом пункта → его код; со склейкой любой пары
    группа + вариант → оба кода. Строка, которой врач заменил выданную строку
    группы, → (код группы, текст врача): «его формулировка этого пункта».
    Совсем своя строка → ('', текст)."""
    import difflib

    rows = list(generated.rows)
    finals = [ln.strip() for ln in final_text.splitlines() if _ws_key(ln)]
    item_by_text: dict[str, Item] = {}
    joined: dict[str, tuple[Item, Item]] = {}
    for it in wb.items:
        item_by_text.setdefault(_ws_key(it.text), it)
        for c in it.children:
            for key in _variant_join_text(it, c):
                joined.setdefault(key, (it, c))

    def by_text(line: str) -> list[tuple[str, str]] | None:
        key = _ws_key(line)
        if key in joined:
            g, c = joined[key]
            return [(g.id, g.text), (c.id, c.text)]
        if key in item_by_text:
            it = item_by_text[key]
            return [(it.id, it.text)]
        return None

    def row_lines(row) -> list[tuple[str, str]]:
        res = []
        for c in row.codes:
            it = wb.by_id.get(c)
            if it is None:
                continue
            if it.children and out and _same_group(out[-1][0], it, wb):
                continue             # второй вариант той же группы — без повтора группы
            res.append((c, it.text))
        return res

    out: list[tuple[str, str]] = []
    sm = difflib.SequenceMatcher(None, [_ws_key(r.text) for r in rows],
                                 [_ws_key(f) for f in finals], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == 'equal':
            for r in rows[i1:i2]:
                out += row_lines(r)
            continue
        replaced = rows[i1:i2] if op == 'replace' else []
        for k, line in enumerate(finals[j1:j2]):
            known = by_text(line)
            if known:
                out += known
            elif k < len(replaced) and len(replaced[k].codes) > 1:
                out.append((replaced[k].codes[0], line))     # своя формулировка пункта
            else:
                out.append(('', line))
    return out
